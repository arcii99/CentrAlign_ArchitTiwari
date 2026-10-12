"""Reading company documents the way a careful clerk would: the PDF's own text when it has one, local OCR when
it is a scan, and spreadsheets as header + rows. Everything runs on the CPU; nothing leaves the machine."""

from __future__ import annotations

import re
from functools import lru_cache
from pathlib import Path
from statistics import median

import pypdfium2 as pdfium
from openpyxl import load_workbook
from pypdf import PdfReader

MIN_TEXT_CHARS = 40  # below this a page is treated as a scan


@lru_cache(maxsize=1)
def _ocr_engine():
    from rapidocr import RapidOCR  # imported lazily: model load takes a second

    return RapidOCR()


def ocr_image(image) -> str:
    """OCR one page image and rebuild its layout: boxes on the same baseline become one line, ' | ' between cells,
    so tables stay readable for a language model."""
    result = _ocr_engine()(image)
    if result.txts is None or not len(result.txts):
        return ""
    items = []
    for box, text in zip(result.boxes, result.txts):
        ys, xs = [p[1] for p in box], [p[0] for p in box]
        items.append(((min(ys) + max(ys)) / 2, min(xs), max(ys) - min(ys), text.strip()))
    tolerance = 0.55 * median(i[2] for i in items)
    lines: list[list] = []
    for item in sorted(items):
        if lines and abs(item[0] - lines[-1][0][0]) <= tolerance:
            lines[-1].append(item)
        else:
            lines.append([item])
    return "\n".join(" | ".join(t for _, _, _, t in sorted(line, key=lambda i: i[1])) for line in lines)


def layout_text(page) -> str:
    """A text page with its layout kept: a table row stays on one line, ' | ' between columns (like ocr_image)."""
    lines: list[str] = []
    for line in (page.extract_text(extraction_mode="layout") or "").splitlines():
        line = re.sub(r" {3,}", " | ", line.strip())
        if line or (lines and lines[-1]):
            lines.append(line)
    return "\n".join(lines).strip()


def read_pdf(path: str | Path, max_chars: int = 12000) -> dict:
    path = Path(path)
    pages, used_ocr = [], False
    reader = PdfReader(path)
    pdf = None
    for n, page in enumerate(reader.pages):
        text = layout_text(page)
        if len(text.replace(" ", "")) < MIN_TEXT_CHARS:
            pdf = pdf or pdfium.PdfDocument(str(path))
            text = ocr_image(pdf[n].render(scale=2).to_pil())
            used_ocr = True
        pages.append(text)
    if pdf:
        pdf.close()
    body = "\n\n".join(f"[page {i + 1}]\n{t}" if len(pages) > 1 else t for i, t in enumerate(pages))
    return {"path": str(path), "pages": len(pages), "ocr": used_ocr, "text": body[:max_chars],
            "truncated": len(body) > max_chars}


def read_xlsx(path: str | Path, sheet: str | None = None, last_rows: int | None = 20) -> dict:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True) if any(v not in (None, "") for v in r)]
    header, body = (rows[0], rows[1:]) if rows else ([], [])
    shown = body[-last_rows:] if last_rows else body
    return {"path": str(path), "sheets": wb.sheetnames, "sheet": ws.title, "header": header,
            "row_count": len(body), "rows_shown": f"last {len(shown)}" if last_rows and len(body) > len(shown) else "all",
            "rows": shown}
