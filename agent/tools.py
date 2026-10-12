"""Nova's tools: the registry the model calls through, and the implementations behind each tool.

Tool hierarchy: the mail server's API, file libraries for the drive, the browser only for the legacy portal.
Each tool returns a Result: compact text for the model, a one-line summary that stays in its history, and
structured data for the event log. Tools never raise; failures come back as ERROR results the model can act on.
"""

from __future__ import annotations

import ast
import base64
import csv
import difflib
import io
import json
import operator
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Any, Callable
from urllib.parse import quote, urlsplit

import httpx
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from agent import config, docs
from agent.human import Pause
from agent.memory import Handbook
from agent.policy import Decision, Policy
from agent.result import Result
from agent.state import Approval, RunState, Step


class ToolError(Exception):
    """A failure the model can fix (bad argument, not found, rejected by a system)."""


# --------------------------------------------------------------------------- context


class MailClient:
    """The mail server's REST API, authenticated with the mailbox's API key from the vault."""

    def __init__(self, url: str, api_key: str, mailbox: str):
        self.mailbox = mailbox
        self.http = httpx.Client(base_url=url, headers={"X-API-Key": api_key}, timeout=30)

    def _call(self, method: str, path: str, **kw) -> Any:
        try:
            r = self.http.request(method, path, **kw)
        except httpx.TransportError as e:
            raise ToolError(f"the mail server is not reachable ({e.__class__.__name__})") from e
        if r.status_code >= 400:
            try:
                detail = r.json().get("detail")
            except ValueError:
                detail = r.text[:200]
            raise ToolError(f"mail server said {r.status_code}: {detail}")
        return r

    def list(self, folder: str, q: str, limit: int = 25) -> list[dict]:
        return self._call("GET", "/api/messages", params={"folder": folder, "q": q, "limit": limit}).json()

    def read(self, mid: int) -> dict:
        return self._call("GET", f"/api/messages/{mid}").json()

    def attachment(self, mid: int, att_id: int) -> bytes:
        return self._call("GET", f"/api/messages/{mid}/attachments/{att_id}").content

    def send(self, payload: dict) -> dict:
        return self._call("POST", "/api/messages", json=payload).json()

    def forward(self, mid: int, payload: dict) -> dict:
        return self._call("POST", f"/api/messages/{mid}/forward", json=payload).json()


@dataclass
class Ctx:
    """Everything a tool may use. Built by the runtime for one run."""
    state: RunState
    drive: Path
    run_dir: Path
    today: date
    handbook: Handbook
    human: Any  # agent.human.Human
    mail: MailClient | None
    sites: dict[str, dict]  # vault entries the browser may sign in to
    log: Callable[..., int]  # log(kind, **data) -> seq
    headless: bool = True
    policy: Policy | None = None  # the gate; None = no checks (only for tests of other parts)
    created_files: set[str] = field(default_factory=set)
    sent: list[dict] = field(default_factory=list)  # emails sent in this run: {"id", "to", "subject"}
    contacts: set[str] = field(default_factory=set)  # vendor contacts seen in the systems of record
    vendors: list[str] = field(default_factory=list)  # "site|vendor ID" this task has come across, oldest first
    rows: list[dict] = field(default_factory=list)  # spreadsheet rows this task added: {"path", "row", "values", ...}
    lessons: Any = None  # agent.memory.Lessons: standing instructions people asked Nova to remember
    done: dict[str, str] = field(default_factory=dict)  # identity of a consequential request that went through → record
    uncertain: set[str] = field(default_factory=set)  # requests answered with a server error, not yet checked
    inflight: dict | None = None  # the consequential request the current browser action sent
    pending_use: str = ""  # approval authorising the action in flight; consumed if the action succeeds
    _browser: Any = None

    def browser(self):
        if self._browser is None:
            from agent.browser import Browser  # Playwright is only loaded when a task needs it

            self._browser = Browser(self.sites, self.run_dir / "downloads", headless=self.headless,
                                    on_post=self.on_post)
        return self._browser

    def site_of(self, url: str) -> str | None:
        origin = "{0.scheme}://{0.netloc}".format(urlsplit(url))
        return next((k for k, v in self.sites.items() if v.get("url", "").rstrip("/") == origin), None)

    def gate(self, action: str, fields: dict, via: str) -> Decision:
        """The policy decision for one consequential request, recorded in the event log."""
        if self.policy is None:
            return Decision("allow")
        if action.startswith("mail."):
            self.lookup_contacts(fields.get("recipients") or [])
        d = self.policy.evaluate(action, fields, self.state.approvals, self.contacts,
                                 context={"correcting": self.state.correcting})
        self.log("policy", action=action, via=via, effect=d.effect, rules=d.rules, approval=d.approval_ref,
                 message=d.message)
        if d.allowed and d.approval_ref:
            self.pending_use = d.approval_ref
        return d

    def on_post(self, req: dict):
        """Every form POST the browser makes passes the policy gate here, at the network layer: whichever clicks
        led to it, a request the policy doesn't allow never reaches the system."""
        site = self.site_of(req.get("url", "")) or "?"
        action, params = config.action_for(site, req.get("method", "POST"), req.get("path", ""))
        self.log("browser.post", site=site, action=action, path=req.get("path"), form=req.get("form"))
        if not action:
            return None
        fields = {**(req.get("form") or {}), **params}
        post = {"site": site, "action": action, "fields": fields, "identity": self.identity(action, fields)}
        if post["identity"] in self.done:  # idempotency: what went through is not sent again
            self.log("runtime.duplicate", action=action, identity=post["identity"], record=self.done[post["identity"]])
            return False, (f"this was already done in this task ({self.done[post['identity']]}): it is not sent "
                           "again. Look at that record instead.")
        if post["identity"] in self.uncertain:  # the last try got a server error: check before anything is resent
            self.inflight = {**post, "recheck": True}
            return False, "the last attempt got an error from the portal and may have gone through, so it is checked first."
        if action in config.UNIQUE:  # a record is entered once: look for it in the system itself first
            verdict, page = self.check_outcome(post)
            self.log("runtime.exists_check", action=action, identity=post["identity"], verdict=verdict, page=page)
            if verdict == "saved":
                return False, (f"this already exists in the system ({page}), so it is not created again. A record is "
                               "entered once: if the procedure says what to do when it is already there (a duplicate, "
                               "an invoice already entered or paid), do that instead.")
        record = str(req.get("path") or "").split("?")[0]  # the record the request changes, e.g. /bills/B-1/status
        d = self.gate(action, {**fields, "record": record}, via="network")
        if d.allowed:
            self.inflight = post
        return None if d.allowed else (False, d.message)

    @staticmethod
    def identity(action: str, fields: dict) -> str:
        """What makes a consequential request "the same request": the fields its confirmation recipe names."""
        page, marker = config.CONFIRM.get(action, ("", ""))
        names = sorted({n.removesuffix("_last4") for n in re.findall(r"\{([a-z_]+)\}", page + marker)})
        return action + "|" + "|".join(f"{n}={str(fields.get(n, '')).strip().lower()}" for n in names) if names else ""

    def check_outcome(self, post: dict) -> tuple[str, str]:
        """("saved" | "not saved" | "unknown", the page read). Read-only, in a separate tab."""
        page, marker = config.CONFIRM.get(post["action"], ("", ""))
        if not page or post["site"] not in self.sites:
            return "unknown", ""
        def value(name: str) -> str:
            if name.endswith("_last4"):
                return str(post["fields"].get(name.removesuffix("_last4"), ""))[-4:]
            return str(post["fields"].get(name, ""))

        path = re.sub(r"\{([a-z_]+)\}", lambda m: quote(value(m[1]), safe="/"), page)
        pattern = re.sub(r"\{([a-z_]+)\}", lambda m: re.escape(value(m[1]).lower().replace("_", " ")), marker)
        url = self.sites[post["site"]]["url"].rstrip("/") + path
        data = self.browser().peek(url, data_only=True)
        if not data:
            return "unknown", path
        # "not saved" only from the record's own page: an error page, or another page, proves nothing (seen
        # 2026-10-09: "not saved" after a 503 that came after the save, and the resend made a second bank account)
        head = data.splitlines()[0]
        record = path.split("?")[0].rstrip("/").split("/")[-1]
        if re.search(r"unavailable|not found|error|sign in", head, re.I) or \
                (re.search(r"\d", record) and record.lower() not in data.lower()):
            return "unknown", path
        hits = [line for line in data.splitlines() if re.search(pattern.lower(), line.lower().replace("_", " "))]
        # the matching line goes to the model: it carries the new record's ID, which later steps need
        return ("saved", f"{path} shows it: {hits[0][:200]}") if hits else ("not saved", path)

    def consume(self, ok: bool, used_by: str) -> None:
        """After the action an approval authorised: if it succeeded, the approval is spent."""
        ref, self.pending_use = self.pending_use, ""
        if ref and ok:
            for a in self.state.approvals:
                if a.ref == ref:
                    a.consumed_by = used_by
                    self.log("approval.used", ref=ref, used_by=used_by)

    def note_contacts(self, url: str, text: str) -> None:
        site = self.site_of(url) or ""
        pattern = config.CONTACTS.get(site)
        if pattern:
            self.contacts.update(m.lower() for m in re.findall(pattern, text))
        if site in config.VENDOR_PAGES:  # from the page's data only: a form's dropdown lists every vendor there is
            data = re.sub(r"^FORM .*(?:\n  .*)*", "", text, flags=re.M)
            for vid in re.findall(config.VENDOR_PAGES[site][0], data):
                key = f"{site}|{vid}"
                if key in self.vendors:
                    self.vendors.remove(key)
                self.vendors.append(key)

    def lookup_contacts(self, recipients: list) -> None:
        """Whether an outside address is a vendor's registered contact is a fact in the vendor master, not something
        Nova has to have looked at: before deciding, read the records of the vendors this task has dealt with (the
        latest three), in a separate tab, read-only. Without this, a run that never opened the vendor's page got its
        remittance held (seen 2026-10-08)."""
        domain = self.policy.domain if self.policy else ""
        unknown = [str(a).strip().lower() for a in recipients if str(a).strip()
                   and not str(a).strip().lower().endswith("@" + domain) and str(a).strip().lower() not in self.contacts]
        if not unknown or not self.vendors:
            return
        read = []
        for key in reversed(self.vendors[-3:]):
            site, vid = key.split("|", 1)
            url = self.sites[site]["url"].rstrip("/") + config.VENDOR_PAGES[site][1].format(id=vid)
            self.note_contacts(url, self.browser().peek(url))
            read.append(vid)
            if all(a in self.contacts for a in unknown):
                break
        self.log("contacts.lookup", addresses=unknown, vendor_records=read,
                 registered=[a for a in unknown if a in self.contacts])

    def close(self) -> None:
        if self._browser is not None:
            self._browser.close()
            self._browser = None


# --------------------------------------------------------------------------- registry


@dataclass
class Tool:
    name: str
    description: str
    params: dict  # JSON-schema properties
    required: list[str]
    fn: Callable[[Ctx, dict], Result]
    kind: str = "read"  # read | write | human | control
    phases: tuple[str, ...] = ("understand", "act")

    def schema(self) -> dict:
        return {"type": "function", "function": {
            "name": self.name, "description": self.description,
            "parameters": {"type": "object", "properties": self.params, "required": self.required}}}


def _plain(n: float | int) -> str:
    return str(int(n)) if float(n).is_integer() else f"{n:.2f}".rstrip("0").rstrip(".")


def _coerce(value: Any, spec: dict) -> Any:
    """Forgive the usual small-model slips: JSON inside strings, numbers as strings, a single item for a list."""
    t = spec.get("type")
    if t == "object" and isinstance(value, str):
        # as written, then with stray escaping undone (seen 2026-10-08: {"Vendor":"V-1001","Invoice no.\":\"GLX..."})
        for text in (value, value.replace('\\"', '"')):
            try:
                parsed = json.loads(text)
                if isinstance(parsed, dict):
                    return parsed
            except ValueError:
                continue
        return value
    if t == "array":
        if isinstance(value, str):
            s = value.strip()
            if s.startswith("["):
                try:
                    return json.loads(s)
                except ValueError:
                    pass
            return [p.strip() for p in re.split(r"[,;]", s) if p.strip()]
        return value if isinstance(value, list) else [value]
    if t == "integer":
        if isinstance(value, str) and re.fullmatch(r"\s*-?\d+\s*", value):
            return int(value)
        if isinstance(value, float) and value.is_integer():
            return int(value)
    if t == "number" and isinstance(value, str):
        try:
            return float(value.replace(",", ""))
        except ValueError:
            return value
    if t == "string" and isinstance(value, (int, float)) and not isinstance(value, bool):
        return _plain(value)
    return value


class Registry:
    def __init__(self, tools: list[Tool]):
        self.tools = {t.name: t for t in tools}

    def schemas(self, phase: str) -> list[dict]:
        return [t.schema() for t in self.tools.values() if phase in t.phases]

    def get(self, name: str) -> Tool | None:
        return self.tools.get(name)

    def call(self, ctx: Ctx, name: str, args: dict, phase: str) -> Result:
        tool = self.tools.get(name)
        if tool is None:
            close = difflib.get_close_matches(name, list(self.tools), n=3)
            return Result.error(f"there is no tool called {name!r}." + (f" Did you mean {', '.join(close)}?" if close else ""))
        if phase not in tool.phases:
            return Result.error(f"{name} is not available yet: first finish understanding the task and call "
                                f"submit_plan." if phase == "understand" else f"{name} is not available now.")
        if not isinstance(args, dict):
            return Result.error(f"arguments for {name} must be an object")
        args = {k: _coerce(v, tool.params[k]) for k, v in args.items() if k in tool.params}
        missing = [k for k in tool.required if args.get(k) in (None, "", [], {})]
        if missing:  # the whole call's shape, with what was given (seen 2026-10-09: "mail_send needs subject." five
            # times, the subject never added)
            shape = ", ".join(f"{k}={json.dumps(args[k], ensure_ascii=False)[:40]}" if k in args and k not in missing
                              else f"{k}=<{tool.params[k].get('description') or k}>" for k in tool.required)
            return Result.error(f"{name} needs {', '.join(missing)}: call {name}({shape}), with every one of them.")
        try:
            return tool.fn(ctx, args)
        except ToolError as e:
            return Result.error(str(e))
        except Pause:  # a person hasn't answered yet: the runtime saves the run and stops here
            raise
        except Exception as e:  # a tool bug must not end the run; the model sees it and the log keeps it
            return Result.error(f"{name} failed unexpectedly: {e.__class__.__name__}: {e}")


# --------------------------------------------------------------------------- handbook


def handbook_search(ctx: Ctx, a: dict) -> Result:
    hits = ctx.handbook.search(a["query"], k=5)
    if not hits:
        return Result(f'No handbook section matches "{a["query"]}". Contents:\n{ctx.handbook.toc()}',
                      f'handbook_search "{a["query"]}": no match')
    lines = [f'Handbook sections matching "{a["query"]}":']
    for s, _ in hits:
        excerpt = re.sub(r"\s+", " ", s.text)[:220]
        lines.append(f"- {s.doc} › {s.heading or s.title}: {excerpt}…")
    lines.append("Read one with handbook_read(document, section).")
    return Result("\n".join(lines), f'handbook_search "{a["query"]}" → ' + "; ".join(
        f"{s.doc} › {s.heading or s.title}" for s, _ in hits[:3]))


def handbook_read(ctx: Ctx, a: dict) -> Result:
    found = ctx.handbook.read(a["document"], a.get("section"))
    if not found:
        raise ToolError(f"no handbook document matches {a['document']!r}. Documents:\n{ctx.handbook.toc()}")
    doc, text = found
    title = ctx.handbook.titles[doc]
    return Result(f"=== Handbook: {doc} ({title}) ===\n{text}", f"Read handbook {doc} ({title})"
                  + (f" › {a['section']}" if a.get("section") else ""), data={"doc": doc})


# --------------------------------------------------------------------------- mail


def _mail(ctx: Ctx) -> MailClient:
    if ctx.mail is None:
        raise ToolError("no mailbox is configured")
    return ctx.mail


def _when(iso: str) -> str:
    try:
        return datetime.fromisoformat(iso).strftime("%d-%m-%Y %H:%M")
    except (TypeError, ValueError):
        return str(iso)


def _addresses(raw: Any) -> list[str]:
    items = raw if isinstance(raw, list) else re.split(r"[,;]", str(raw or ""))
    out = [str(a).strip().strip("<>").strip() for a in items if str(a).strip()]
    bad = [x for x in out if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", x)]
    if bad:
        raise ToolError(f"not an email address: {', '.join(bad)}")
    return out


def mail_list(ctx: Ctx, a: dict) -> Result:
    folder, search = a.get("folder") or "inbox", a.get("search") or ""
    msgs = _mail(ctx).list(folder, search)
    head = f'{folder}: {len(msgs)} message(s)' + (f' matching "{search}"' if search else "") + " (newest first)"
    if not msgs:
        return Result(head, head)
    lines = [head] + [
        f'#{m["id"]} {_when(m["date"])} from {m["from"]} to {m["to"]}: "{m["subject"]}"'
        + (f' [{m["attachment_count"]} attachment(s)]' if m["attachment_count"] else "")
        + ("" if m["is_read"] else " (unread)") for m in msgs]
    return Result("\n".join(lines), head + ": " + "; ".join(f'#{m["id"]} "{m["subject"][:60]}"' for m in msgs[:4]))


def mail_read(ctx: Ctx, a: dict) -> Result:
    m = _mail(ctx).read(int(a["id"]))
    atts = "; ".join(f'{x["id"]}: {x["filename"]}' for x in m["attachments"]) or "none"
    text = (f'Message #{m["id"]} ({m["folder"]})\nFrom: {m["from"]}\nTo: {m["to"]}\n' + (f'Cc: {m["cc"]}\n' if m["cc"] else "")
            + f'Date: {_when(m["date"])}\nSubject: {m["subject"]}\nAttachments: {atts}\n\n{m["body"]}')
    return Result(text, f'Read message #{m["id"]} from {m["from"]}: "{m["subject"]}" (attachments: {atts})',
                  data={"id": m["id"], "from": m["from"], "subject": m["subject"]})


def _read_file(path: Path, sheet: str | None = None, rows: int = 10) -> tuple[str, str]:
    """(text, kind) for a document on disk."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        r = docs.read_pdf(path)
        kind = f"PDF, {r['pages']} page(s), " + ("scanned: read by OCR" if r["ocr"] else "text")
        return r["text"] + ("\n[… truncated]" if r["truncated"] else ""), kind
    if ext in (".xlsx", ".xlsm"):
        return _sheet_text(path, sheet, rows), "spreadsheet"
    if ext == ".csv":
        lines = path.read_text(encoding="utf-8-sig", errors="replace").splitlines()
        body = "\n".join(lines[:40]) + (f"\n[… {len(lines) - 40} more lines]" if len(lines) > 40 else "")
        return body, f"CSV, {max(len(lines) - 1, 0)} data rows"
    if ext in (".md", ".txt", ".json", ".html", ".eml"):
        text = path.read_text(encoding="utf-8", errors="replace")
        return text[:8000] + ("\n[… truncated]" if len(text) > 8000 else ""), "text"
    return f"(binary file, {path.stat().st_size} bytes; cannot be shown as text)", "binary"


def mail_attachment(ctx: Ctx, a: dict) -> Result:
    mid = int(a["message_id"])
    m = _mail(ctx).read(mid)
    want = re.sub(r"^\d+\s*:\s*", "", str(a["attachment"]).strip()) or str(a["attachment"]).strip()  # "6: x.pdf"
    atts = m["attachments"]
    match = [x for x in atts if str(x["id"]) == want] or [x for x in atts if x["filename"].lower() == want.lower()] \
        or [x for x in atts if want.lower() in x["filename"].lower()]
    if len(match) != 1:
        listing = "; ".join(f'{x["id"]}: {x["filename"]}' for x in atts) or "none"
        raise ToolError(f"message #{mid} has no single attachment matching {want!r}. Its attachments: {listing}")
    att = match[0]
    folder = ctx.run_dir / "attachments"
    folder.mkdir(parents=True, exist_ok=True)
    local = folder / f"{mid}-{Path(att['filename']).name}"
    local.write_bytes(_mail(ctx).attachment(mid, att["id"]))
    text, kind = _read_file(local)
    rel = f"attachments/{local.name}"
    return Result(f"Attachment {att['filename']} of message #{mid} ({kind}), saved as {rel}:\n{text}",
                  f"Read attachment {att['filename']} of message #{mid} ({kind})", data={"path": rel})


_CLOSING = re.compile(r"^(best |kind |warm |with )?(regards|thanks|thank you|many thanks|best|sincerely|yours "
                      r"sincerely|yours faithfully|cheers)[,.!]?$", re.I)


def _sign(ctx: Ctx, body: str) -> str:
    """The body as it goes out, signed as the Email Guidelines say. An email signed in someone else's name is refused
    (seen 2026-10-09: a refund request to Acme signed "Meera Joshi, Accounts, Acme Industrial Supplies", the name on
    Acme's own complaint; another signed "AP Team"). Without a sign-off, the company's is added."""
    sig = str((ctx.policy.company if ctx.policy else {}).get("email_signature") or "").strip()
    if not sig:
        return body
    lines = [x.strip() for x in str(body).strip().splitlines() if x.strip()]
    me = sig.splitlines()[0].split("(")[0].strip().lower()  # "nova"
    if any(me in x.lower() for x in lines[-4:]):
        return body
    closing = next((i for i in range(len(lines) - 1, max(-1, len(lines) - 7), -1) if _CLOSING.match(lines[i])), None)
    if closing is not None and closing < len(lines) - 1:
        raise ToolError(f"not sent: it is signed as \"{' / '.join(lines[closing + 1:])}\", but Nova writes in its own "
                        f"name, from Northwind's AP mailbox. Sign off as the Email Guidelines say:\n{sig}\n"
                        "Read the body again too: it must speak for Northwind, not for the person you are writing to.")
    return f"{str(body).rstrip()}\n\n{sig}"


def mail_send(ctx: Ctx, a: dict) -> Result:
    to, cc = _addresses(a["to"]), _addresses(a.get("cc") or [])
    body = _sign(ctx, a["body"])
    who, subject = sorted(x.lower() for x in to + cc), re.sub(r"\s+", " ", str(a["subject"])).strip().lower()
    for prior in ctx.sent:  # sending is not idempotent, so a retry must not send it twice
        if prior["who"] == who and prior["subject"] == subject:
            raise ToolError(f"this email was already sent in this task (#{prior['id']}, same recipients and "
                            "subject), so it was not sent again. If that step is finished, mark it with step_done.")
    files = []
    for p in a.get("attachments") or []:
        path = _resolve(ctx, p)
        files.append({"filename": path.name, "content_base64": base64.b64encode(path.read_bytes()).decode()})
    d = ctx.gate("mail.send", {"to": to, "cc": cc, "recipients": to + cc, "subject": a["subject"], "body": a["body"]},
                 via="tool")
    if not d.allowed:
        raise ToolError(f"blocked by policy: {d.message}")
    payload = {"to": to, "cc": cc, "subject": a["subject"], "body": body, "attachments": files}
    try:
        out = _mail(ctx).send(payload)
    except ToolError:
        ctx.consume(False, "")
        raise
    ctx.sent.append({"id": out["id"], "who": who, "subject": subject})
    ctx.consume(True, f"message #{out['id']}")
    ctx.log("mail.sent", id=out["id"], to=to, cc=cc, subject=a["subject"], body=body,
            attachments=[f["filename"] for f in files])
    return Result(f'Sent message #{out["id"]} to {", ".join(to)}' + (f' (cc {", ".join(cc)})' if cc else "")
                  + f': "{a["subject"]}"' + (f' with {", ".join(f["filename"] for f in files)}' if files else ""),
                  f'Sent #{out["id"]} to {", ".join(to)}: "{a["subject"]}"', data={"id": out["id"]})


def mail_forward(ctx: Ctx, a: dict) -> Result:
    to, cc = _addresses(a["to"]), _addresses(a.get("cc") or [])
    try:  # what goes out is the original under "Fwd:", plus the note: an approval binds both
        subject = "Fwd: " + str(_mail(ctx).read(int(a["message_id"])).get("subject") or "")
    except (ToolError, ValueError, TypeError):
        subject = "Fwd:"
    d = ctx.gate("mail.forward", {"to": to, "cc": cc, "recipients": to + cc, "subject": subject,
                                  "body": a.get("body") or ""}, via="tool")
    if not d.allowed:
        raise ToolError(f"blocked by policy: {d.message}")
    try:
        out = _mail(ctx).forward(int(a["message_id"]), {"to": to, "cc": cc, "body": a.get("body") or ""})
    except ToolError:
        ctx.consume(False, "")
        raise
    ctx.consume(True, f"message #{out['id']}")
    ctx.log("mail.forwarded", id=out["id"], original=int(a["message_id"]), to=to, cc=cc, body=a.get("body"))
    return Result(f'Forwarded message #{a["message_id"]} (with its attachments) to {", ".join(to)} as #{out["id"]}',
                  f'Forwarded #{a["message_id"]} to {", ".join(to)}', data={"id": out["id"]})


# --------------------------------------------------------------------------- drive

RUN_FOLDERS = ("attachments", "downloads")  # the run's own files, addressable like drive folders


def _resolve(ctx: Ctx, raw: str, must_exist: bool = True) -> Path:
    return resolve_path(ctx.drive, ctx.run_dir, raw, must_exist)


def resolve_path(drive: Path, run_dir: Path, raw: str, must_exist: bool = True) -> Path:
    """A path on the shared drive (or in the run's own folders), however it was written; never outside them."""
    p = str(raw).strip().strip('"').replace("\\", "/")
    if re.match(r"^[a-zA-Z]:", p) or p.startswith("//"):
        raise ToolError("use a path on the shared drive, relative to its root (e.g. AP/AP_Tracker.xlsx)")
    p = re.sub(r"^/+", "", p)
    p = re.sub(r"^(drive|shared drive)/", "", p, flags=re.I)
    base = run_dir if p.split("/")[0] in RUN_FOLDERS else drive
    full = (base / p).resolve()
    if not full.is_relative_to(base.resolve()):
        raise ToolError("that path is outside the shared drive")
    if must_exist and not full.exists():
        parent = full.parent
        hint = ""
        if parent.exists() and parent.is_dir():
            near = difflib.get_close_matches(full.name, [c.name for c in parent.iterdir()], n=3, cutoff=0.4)
            hint = f" Similar names in that folder: {', '.join(near)}" if near else ""
        raise ToolError(f"{raw} does not exist.{hint}")
    return full


def _rel(ctx: Ctx, path: Path) -> str:
    for base in (ctx.drive, ctx.run_dir):
        if path.resolve().is_relative_to(base.resolve()):
            return path.resolve().relative_to(base.resolve()).as_posix()
    return path.name


def files_list(ctx: Ctx, a: dict) -> Result:
    folder = _resolve(ctx, a.get("folder") or ".")
    if not folder.is_dir():
        raise ToolError(f"{a.get('folder')} is a file, not a folder; read it with file_read")
    entries = sorted(folder.iterdir(), key=lambda p: (p.is_file(), p.name.lower()))
    lines = [f"{_rel(ctx, folder) if folder != ctx.drive.resolve() else '(drive root)'}:"]
    for p in entries:
        if p.is_dir():
            lines.append(f"  {p.name}/")
        else:
            st = p.stat()
            lines.append(f"  {p.name}  ({st.st_size // 1024 or 1} KB, modified "
                         f"{datetime.fromtimestamp(st.st_mtime):%d-%m-%Y})")
    if folder == ctx.drive.resolve():
        lines.append("  (this run's own files: attachments/, downloads/)")
    return Result("\n".join(lines), f"Listed {a.get('folder') or 'drive root'}: {len(entries)} item(s)")


def _fmt(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        return _plain(v)
    if isinstance(v, datetime):
        return v.strftime("%d-%m-%Y")
    return str(v)


def _sheet_text(path: Path, sheet: str | None, rows: int) -> str:
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        if sheet and sheet not in wb.sheetnames:
            raise ToolError(f"no sheet {sheet!r}; sheets: {', '.join(wb.sheetnames)}")
        ws = wb[sheet] if sheet else wb.worksheets[0]
        data = [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    numbered = [(i, r) for i, r in enumerate(data, 1) if any(v not in (None, "") for v in r)]
    if not numbered:
        return f'Sheet "{ws.title}" is empty.'
    (hrow, header), body = numbered[0], numbered[1:]
    width = max(i for i, h in enumerate(header, 1) if h not in (None, ""))
    shown = body[-rows:] if rows else body
    lines = [f'Sheet "{ws.title}" (sheets: {", ".join(wb.sheetnames)}): {len(body)} data rows'
             + (f", showing the last {len(shown)}" if len(shown) < len(body) else ""),
             "Columns: " + " | ".join(_fmt(h) for h in header[:width])]
    lines += [f"row {i}: " + " | ".join(_fmt(v) for v in r[:width]) for i, r in shown]
    return "\n".join(lines)


def file_read(ctx: Ctx, a: dict) -> Result:
    path = _resolve(ctx, a["path"])
    if path.is_dir():
        return files_list(ctx, {"folder": a["path"]})
    text, kind = _read_file(path, a.get("sheet"), min(int(a.get("rows") or 10), 25))
    return Result(f"{_rel(ctx, path)} ({kind}):\n{text}", f"Read {_rel(ctx, path)} ({kind})")


def _cell_value(value: Any, existing: list[Any]) -> Any:
    """Numbers stay numbers in columns that hold numbers ('1,74,000.00' → 174000); everything else is text."""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    s = "" if value is None else str(value).strip()
    numeric_col = existing and all(isinstance(v, (int, float)) for v in existing)
    t = re.sub(r"^(₹|rs\.?|inr)\s*", "", s.replace(",", ""), flags=re.I)
    if numeric_col and re.fullmatch(r"-?\d+(\.\d+)?", t):
        return float(t) if "." in t else int(t)
    return s or None


def sheet_append_row(ctx: Ctx, a: dict) -> Result:
    path = _resolve(ctx, a["path"])
    row = a["row"]
    if not isinstance(row, dict) or not row:
        raise ToolError('row must be an object of {"column name": value}')
    try:
        wb = load_workbook(path)
    except PermissionError as e:
        raise ToolError(f"{a['path']} is locked (open in another program?)") from e
    ws = wb[a["sheet"]] if a.get("sheet") in wb.sheetnames else wb.worksheets[0]
    header = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    norm = lambda s: re.sub(r"[^a-z0-9]+", "", str(s).lower())  # noqa: E731
    cols = {norm(h): i for i, h in enumerate(header) if h}
    unknown = [k for k in row if norm(k) not in cols]
    if unknown:
        raise ToolError(f"unknown column(s) {', '.join(map(repr, unknown))}. The columns are: {' | '.join(h for h in header if h)}")
    existing_rows = [list(r) for r in ws.iter_rows(min_row=2, values_only=True) if any(v not in (None, "") for v in r)]
    values: list[Any] = [None] * len(header)
    for k, v in row.items():
        i = cols[norm(k)]
        recent = [r[i] for r in existing_rows[-30:] if i < len(r) and r[i] not in (None, "")]
        values[i] = _cell_value(v, recent)
    given = [i for i, v in enumerate(values) if v not in (None, "")]
    for n, r in enumerate(existing_rows, 2):
        if given and all(i < len(r) and _fmt(r[i]) == _fmt(values[i]) for i in given):
            raise ToolError(f"an identical row already exists (row {n}); it was not added again")
    rel = _rel(ctx, path)
    _check_row(ctx, rel, header, values, existing_rows, a, norm)
    d = ctx.gate("file.append", {**{h: _fmt(v) for h, v in zip(header, values) if h}, "path": rel, "record": rel},
                 via="tool")
    if not d.allowed:
        raise ToolError(f"blocked by policy: {d.message}")
    ws.append(values)
    try:
        wb.save(path)
    except PermissionError as e:
        ctx.consume(False, "")
        raise ToolError(f"{a['path']} is locked (open in another program?); the row was not saved") from e
    ctx.consume(True, f"{rel} row {ws.max_row}")
    try:
        fix = int(a["correction"]) if a.get("correction") not in (None, "") else None
    except (TypeError, ValueError):
        fix = None
    ctx.rows.append({"path": rel, "row": ws.max_row, "values": [_fmt(v) for v in values], "correction": fix})
    blank = [h for h, v in zip(header, values) if h and v in (None, "")]
    shown = " | ".join(f"{h}={_fmt(v)}" for h, v in zip(header, values) if h and v not in (None, ""))
    ctx.log("file.row_appended", path=_rel(ctx, path), row=ws.max_row, values=dict(zip(header, map(_fmt, values))))
    return Result(f"Added row {ws.max_row} to {_rel(ctx, path)} ({ws.title}): {shown}"
                  + (f"\nLeft blank: {', '.join(blank)}" if blank else ""),
                  f"Added row {ws.max_row} to {_rel(ctx, path)}: {shown[:160]}")


def _check_row(ctx: Ctx, rel: str, header: list[str], values: list, history: list[list], a: dict, norm) -> None:
    """Two rules a log like the AP tracker lives by, read from the sheet's own history (seen 2026-10-07/08: a row
    without its Status, then "fixed" by appending more rows; five rows for one invoice).
    1. A column filled in on every earlier row is not left empty by accident: fill it, or confirm with `blank`.
    2. One row per document per task: a second row for the same invoice is only a single, explained correction."""
    if len(history) < 5:
        return
    filled = lambda i: [r[i] for r in history if i < len(r) and r[i] not in (None, "")]  # noqa: E731
    usual = [i for i, h in enumerate(header) if h and len(filled(i)) >= 0.9 * len(history)]
    confirmed = {norm(b) for b in (a.get("blank") or []) if str(b).strip()}
    empty = [header[i] for i in usual if values[i] in (None, "") and norm(header[i]) not in confirmed]
    if empty:
        # where to look, not just "fill them in" (seen 2026-10-09: the same incomplete row sent five times, though
        # the bill page showing its PO had been opened)
        raise ToolError(f"not added: these columns are filled in on every earlier row but are empty here: "
                        f"{', '.join(empty)}. Fill them in: each is on the document or the record this row is about "
                        "(open it and copy the value). Only where the sheet's guide allows a column to be blank "
                        f"(it says \"or blank\", e.g. no bill exists yet) may you confirm it with "
                        f"blank={json.dumps(empty)}.")
    # a value of the wrong kind: where every earlier value in a column has one shape (PO-9999-9999, 99-99-9999), a new
    # value must have it too (seen 2026-10-08: "GSTIN: 29AAFCN5678J1Z4" written as a PO number, to fill the column).
    # Only for columns whose values carry digits: a status like "Payment Scheduled" isn't judged by history's "Paid"
    shape = lambda v: re.sub(r"[A-Za-z]", "A", re.sub(r"\d", "9", _fmt(v).strip()))  # noqa: E731
    for i in usual:
        earlier = {shape(v) for v in filled(i)}
        if values[i] in (None, "") or len(earlier) != 1 or not re.search(r"\d", next(iter(earlier))):
            continue
        if shape(values[i]) not in earlier:
            example = _fmt(filled(i)[-1])
            raise ToolError(f"not added: {header[i]} looks like {example} on every earlier row, and "
                            f"{_fmt(values[i])!r} doesn't. Use the value from the document in that form; if there is "
                            f"none and the sheet's guide allows the column to be blank, leave it out and confirm it "
                            f"with blank={json.dumps([header[i]])}.")
    # the document's own identifier: the first usual column whose earlier values look like IDs and are (nearly) all
    # different. Nearly: a real log holds its own mistakes (this tracker lists one Acme invoice twice)
    key = next((i for i in usual if len(set(map(_fmt, filled(i)))) >= 0.9 * len(filled(i))
                and sum(bool(re.search(r"[A-Za-z]", _fmt(v)) and re.search(r"\d", _fmt(v))) for v in filled(i))
                >= 0.9 * len(filled(i))), None)
    if key is None:
        return
    # a log of documents gets no row without one (seen 2026-10-08: a vendor set-up logged as an invoice "On Hold",
    # every invoice column confirmed blank)
    if values[key] in (None, ""):
        raise ToolError(f"not added: every row of this sheet records one document by its {header[key]}, and this row "
                        f"has none. Work that isn't about such a document (e.g. setting up a vendor) gets no row here: "
                        f"if the procedure doesn't say to log it, skip this step.")
    mine =[r for r in ctx.rows if r["path"] == rel and key < len(r["values"]) and r["values"][key] == _fmt(values[key])]
    if not mine:
        return
    first, fix = mine[0]["row"], a.get("correction")
    try:
        fix = int(fix) if fix not in (None, "") else None
    except (TypeError, ValueError):
        fix = None
    if fix is not None and fix not in {r["row"] for r in mine}:
        raise ToolError(f"not added: correction must be the number of the row it corrects (row {first}).")
    if fix is None:
        raise ToolError(f"not added: row {first}, added earlier in this task, already records {header[key]} "
                        f"{_fmt(values[key])}. The sheet is append-only, one row per document. If row {first} is "
                        f"wrong, add ONE correction row: the whole row as it should be, Notes saying what it "
                        f"corrects, and correction={first}.")
    fixed = [r["row"] for r in mine if r.get("correction") == fix]
    if fixed:
        raise ToolError(f"not added: row {fix} was already corrected by row {fixed[0]} in this task.")
    notes = [i for i, h in enumerate(header) if norm(h) in ("notes", "note", "remarks", "comments")]
    if notes and values[notes[0]] in (None, ""):
        raise ToolError(f"not added: a correction row needs {header[notes[0]]} saying what it corrects in row {fix}.")


def sheet_write(ctx: Ctx, a: dict) -> Result:
    path = _resolve(ctx, a["path"], must_exist=False)
    if path.suffix.lower() != ".xlsx":
        raise ToolError("the file name must end in .xlsx")
    if path.exists() and _rel(ctx, path) not in ctx.created_files:
        raise ToolError(f"{a['path']} already exists; existing files are never overwritten. Choose a new name.")
    columns, rows = [str(c) for c in a["columns"]], a.get("rows") or []
    if any(not isinstance(r, list) or len(r) != len(columns) for r in rows):
        raise ToolError(f"every row must be a list of {len(columns)} values, one per column")
    wb = Workbook()
    ws = wb.active
    ws.title = (a.get("sheet") or "Sheet1")[:31]
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([_cell_value(v, [0]) for v in r])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    ctx.created_files.add(_rel(ctx, path))
    ctx.log("file.written", path=_rel(ctx, path), sheet=ws.title, rows=len(rows))
    return Result(f"Created {_rel(ctx, path)} with sheet {ws.title!r}: {len(rows)} row(s) under "
                  f"{' | '.join(columns)}", f"Created {_rel(ctx, path)} ({len(rows)} rows)")


# --------------------------------------------------------------------------- browser


def _after_browser(ctx: Ctx, res: Result) -> Result:
    """What the runtime learns from every page: whether a submission went through (checked in the system itself when
    the server answered with an error), an approval spent by it, vendor contacts."""
    url = str(res.data.get("url", "")).split("?")[0]
    post, ctx.inflight = ctx.inflight, None
    if post and (res.data.get("post_error") or post.get("recheck")):
        res = _settle_outcome(ctx, post, res)
    else:
        ctx.consume(res.ok, url)
        if post and res.ok and post["identity"]:
            ctx.done[post["identity"]] = url
    ctx.note_contacts(str(res.data.get("url", "")), res.text)
    return res


def _settle_outcome(ctx: Ctx, post: dict, res: Result) -> Result:
    """The portal failed a consequential submission, or one is being resent after such a failure: look in the system
    whether it went through, before anything else happens. Never a blind retry (that is how invoices get paid twice)."""
    verdict, page = ctx.check_outcome(post)
    ctx.log("outcome.check", action=post["action"], identity=post["identity"], page=page, verdict=verdict)
    if verdict == "saved":
        where = page.split(" shows it")[0]
        ctx.consume(True, where)
        ctx.done[post["identity"]] = where
        ctx.uncertain.discard(post["identity"])
        note = f"CHECKED: despite the error, the request went through: {page}. It is done; do not submit it again."
        return Result(f"{res.text}\n\n{note}", f"{res.summary} — but it went through: {page}"[:300], ok=True,
                      data=res.data)
    ctx.consume(False, "")
    if verdict == "not saved":
        ctx.uncertain.discard(post["identity"])
        form = getattr(ctx.browser(), "form_page", "")
        note = (f"CHECKED: the request did NOT go through ({page} does not show it). It is safe to submit it again"
                + (f": this page is the error, not the form, so open {form} again (browser_open), fill it in and "
                   "submit" if form else "") + ".")  # (seen 2026-10-09: refilled twice on the error page itself)
    else:
        ctx.uncertain.add(post["identity"])
        note = (f"Whether it went through could not be checked. Open {page or 'the record'} and look before "
                "submitting again; until then the same request is held.")
    return Result(f"{res.text}\n\n{note}", f"{res.summary} — {note[:120]}", ok=False, data=res.data)


def browser_open(ctx: Ctx, a: dict) -> Result:
    return _after_browser(ctx, ctx.browser().open(str(a["url"])))


def browser_fill(ctx: Ctx, a: dict) -> Result:
    fields = a.get("fields") or {}
    if not isinstance(fields, dict):
        raise ToolError('fields must be an object of {"field label": value}')
    return _after_browser(ctx, ctx.browser().fill(fields, a.get("submit") or None))


def browser_click(ctx: Ctx, a: dict) -> Result:
    return _after_browser(ctx, ctx.browser().click(str(a["target"])))


# --------------------------------------------------------------------------- calculation

_OPS = {ast.Add: operator.add, ast.Sub: operator.sub, ast.Mult: operator.mul, ast.Div: operator.truediv,
        ast.FloorDiv: operator.floordiv, ast.Mod: operator.mod}


def _round(x: Decimal, nd: Decimal = Decimal(0)) -> Decimal:
    return x.quantize(Decimal(1).scaleb(-int(nd)), rounding=ROUND_HALF_UP)


_FUNCS = {"round": _round, "min": min, "max": max, "abs": abs}


def _eval(node: ast.AST) -> Decimal:
    if isinstance(node, ast.Expression):
        return _eval(node.body)
    if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)) and not isinstance(node.value, bool):
        return Decimal(str(node.value))
    if isinstance(node, ast.BinOp) and type(node.op) in _OPS:
        return _OPS[type(node.op)](_eval(node.left), _eval(node.right))
    if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.USub, ast.UAdd)):
        v = _eval(node.operand)
        return -v if isinstance(node.op, ast.USub) else v
    if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and node.func.id in _FUNCS and not node.keywords:
        return _FUNCS[node.func.id](*[_eval(x) for x in node.args])
    raise ToolError("only numbers, + - * / ( ) and round/min/max/abs are allowed")


def calc(ctx: Ctx, a: dict) -> Result:
    expr = str(a["expression"])
    clean = re.sub(r"(₹|\bRs\.?|\bINR\b)", "", expr, flags=re.I)
    clean = re.sub(r"(?<=\d),(?=\d)", "", clean)
    clean = re.sub(r"(?<=[\d)])\s*[x×]\s*(?=[\d(])", " * ", clean)
    clean = re.sub(r"(\d+(?:\.\d+)?)\s*%", r"(\1/100)", clean)
    try:
        value = _eval(ast.parse(clean.strip(), mode="eval"))
    except SyntaxError as e:
        raise ToolError(f"cannot read the expression {expr!r}") from e
    except (ZeroDivisionError, ArithmeticError) as e:
        raise ToolError(f"arithmetic error in {expr!r}: {e}") from e
    v = _round(value, Decimal(2))
    out = str(int(v)) if v == v.to_integral_value() else f"{v:.2f}"
    return Result(f"{expr} = {out}", f"calc {expr} = {out}", data={"value": out})


_DATE_FORMATS = ("%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%Y-%m-%d", "%d-%b-%Y", "%d %b %Y", "%d-%B-%Y", "%d %B %Y",
                 "%B %d, %Y", "%b %d, %Y", "%d %b, %Y", "%d/%m/%y", "%d-%m-%y")


def parse_date(raw: str, today: date) -> date:
    s = re.sub(r"\s+", " ", str(raw).strip())
    if s.lower() == "today":
        return today
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ToolError(f"cannot read the date {raw!r}; write it like 30-09-2026")


def _relation(d: date, today: date) -> str:
    n = (d - today).days
    return "today" if n == 0 else f"{abs(n)} day(s) {'after' if n > 0 else 'before'} today ({today:%d-%m-%Y})"


def date_calc(ctx: Ctx, a: dict) -> Result:
    d = parse_date(a["date"], ctx.today)
    days = int(a.get("add_days") or 0)
    out = d + timedelta(days=days)
    text = (f"{d:%d-%m-%Y} + {days} days = {out:%d-%m-%Y} ({out:%A}), {_relation(out, ctx.today)}" if days
            else f"{a['date']} = {out:%d-%m-%Y} ({out:%A}), {_relation(out, ctx.today)}")
    return Result(text, text, data={"date": f"{out:%d-%m-%Y}"})


# --------------------------------------------------------------------------- people


def _same_question(q: str, earlier: str) -> bool:
    """The same question asked again in other words? (Seen 2026-10-08: one account mismatch asked about three times.)"""
    words = lambda s: {w for w in re.findall(r"[a-z0-9]+", s.lower()) if len(w) > 2}  # noqa: E731
    a, b = words(q), words(earlier)
    overlap = len(a & b) / max(1, len(a | b))
    ids = set(re.findall(r"\d{4,}", q))
    return overlap >= 0.5 or (bool(ids) and ids <= set(re.findall(r"\d{4,}", earlier)) and overlap >= 0.3)


def ask_human(ctx: Ctx, a: dict) -> Result:
    q = str(a["question"]).strip()
    prior = next((p for p in ctx.state.questions if _same_question(q, p["question"])), None)
    if prior and ctx.state.waiting_for == "":  # a person is not interrupted twice for one thing
        ctx.log("human.repeat", question=q, earlier=prior["question"], answer=prior["answer"])
        return Result(f"You asked this already, and {prior.get('by', 'they')} answered: {prior['answer']} "
                      "Act on that answer.", f'Already answered: "{prior["answer"][:800]}"')
    key = f"Q-{ctx.state.run_id[-4:].upper()}-{len(ctx.state.questions) + 1}"
    if ctx.state.waiting_for != key:  # asked once; after a pause the same call comes back for the answer
        ctx.log("human.question", key=key, question=q, to=ctx.state.requester)
    ctx.state.waiting_for = key
    who, answer = ctx.human.ask(q, ctx.state.requester, key, ctx.state.run_id)  # may pause the run
    ctx.state.waiting_for = ""
    ctx.state.questions.append({"question": q, "answer": answer, "by": who})
    ctx.log("human.answer", key=key, by=who, answer=answer)
    _learn(ctx, key, "answer", q, answer, who)
    # the answer is kept whole in the one line that stays in view (seen 2026-10-09: cut at 120 characters, it read
    # as "Use the" and the model asked the person to repeat themselves)
    return Result(f"Answer from {who}: {answer}", f'Asked "{q[:90]}" → {who}: "{answer[:800]}"')


def _learn(ctx: Ctx, key: str, kind: str, situation: str, guidance: str, by: str) -> None:
    """LEARN: an answer, or a rejection's reason, that the person asked Nova to remember becomes a standing
    instruction (agent/memory.py). Only people's own words, and only when they say so."""
    if ctx.lessons is None or not guidance.strip() or not getattr(ctx.human, "remembered", {}).get(key):
        return
    lesson = ctx.lessons.add(kind, situation, guidance, by, ctx.state.run_id, ctx.state.request)
    ctx.log("lesson.learned", id=lesson.id, kind=kind, situation=situation, guidance=guidance, by=by)


_AMOUNT_KEY = re.compile(r"amount|total|value|net|tds|gst|sum", re.I)


def _normalise_payload(details: dict) -> dict:
    """Amounts become numbers ('₹1,74,000.00' → 174000.0), so an approval can later be matched exactly against
    what is actually submitted. Identifiers stay text."""
    out = {}
    for k, v in details.items():
        if isinstance(v, str):
            m = re.fullmatch(r"(₹|rs\.?|inr)?\s*(\d[\d,]*(?:\.\d+)?)", v.strip(), re.I)
            if m and (m[1] or "," in m[2] or _AMOUNT_KEY.search(str(k))):
                v = float(m[2].replace(",", ""))
        out[str(k)] = v
    return out


def request_approval(ctx: Ctx, a: dict) -> Result:
    kinds = ctx.policy.approvals if ctx.policy else {}
    kind = str(a["action"]).strip()
    if kind not in kinds:
        raise ToolError(f"action must be one of: {', '.join(kinds) or '(no approvals are configured)'}")
    details = a["details"]
    if not isinstance(details, dict) or not details:
        raise ToolError("details must be an object saying exactly what is approved (e.g. vendor, bill, amount)")
    details = _normalise_payload(details)
    missing = [k for k in ctx.policy.required_details(kind) if details.get(k) in (None, "", [])]
    if missing:
        raise ToolError(f"details must include {', '.join(missing)}: an approval covers exactly those, and the action "
                        "it approves must match them")
    for prior in ctx.state.approvals:
        if prior.kind == kind and prior.payload == details and prior.status == "approved" and not prior.consumed_by:
            return Result(f"This exact action was already approved: ref {prior.ref}.",
                          f"Approval {prior.ref} already granted for {kind}", data={"ref": prior.ref})
    spec = kinds[kind]
    ref = f"APR-{ctx.state.run_id[-4:].upper()}-{len(ctx.state.approvals) + 1}"
    approver, reason = spec["approver"], str(a.get("reason") or "").strip()  # the policy says who approves
    if ctx.state.waiting_for != ref:
        ctx.log("approval.request", ref=ref, approver=approver, asked=str(a.get("approver") or ""),
                action=spec["action"], kind=kind, payload=details, reason=reason)
    ctx.state.waiting_for = ref
    ok, by, comment = ctx.human.approve(approver, kind, details, reason, ref, ctx.state.run_id)  # may pause the run
    ctx.state.waiting_for = ""
    ctx.state.approvals.append(Approval(ref=ref, approver=approver, action=spec["action"], kind=kind, payload=details,
                                        status="approved" if ok else "rejected", by=by, comment=comment))
    ctx.log("approval.decision", ref=ref, status="approved" if ok else "rejected", by=by, comment=comment)
    shown = ", ".join(f"{k}={v}" for k, v in details.items())
    if not ok:
        _learn(ctx, ref, "rejection", f"approval asked for {kind}: {shown}", comment, by)
    if ok:
        where = f" Enter {ref} in the '{spec['ref_label']}' field." if spec.get("ref_field") else ""
        return Result(f"APPROVED by {by}. Approval ref: {ref}. It covers exactly: {shown}, once.{where}"
                      + (f" Comment: {comment}" if comment else ""),
                      f"Approval {ref} GRANTED by {by} for {kind}: {shown}", data={"ref": ref})
    return Result(f"REJECTED by {by}: {comment or 'no reason given'}. Do not carry out this action; tell the "
                  "requester (ask_human) or finish with an explanation.",
                  f"Approval for {kind} REJECTED by {by}: {comment[:100]}", ok=False, data={"ref": ref})


# --------------------------------------------------------------------------- plan control


def _clean_step(s: Any) -> str:
    return re.sub(r"^\s*(step\s*)?\d+\s*[.):—–-]*\s*", "", str(s), flags=re.I).strip()


def plan_text(state: RunState, width: int = 200) -> str:
    lines = []
    for i, s in enumerate(state.plan, 1):
        mark = {"done": "[done]", "skipped": "[skipped]"}.get(s.status, "[ ]")
        note = re.sub(r"\s+", " ", s.result)
        lines.append(f"{i}. {mark} {s.text}" + (f" → {note[:width]}{'…' if len(note) > width else ''}" if note else ""))
    return "\n".join(lines)


def split_work(ctx: Ctx, a: dict) -> Result:
    """A request that covers several separate documents (an inbox of invoices) becomes one run per document, done one
    after another, each with its own contract, plan and verification: how a person works through an inbox."""
    st = ctx.state
    if st.parent:
        raise ToolError("this request is already one item of a larger one: carry it out (read its documents, then "
                        "submit_plan)")
    items = [str(i.get("request", "") if isinstance(i, dict) else i).strip() for i in a.get("items") or []]
    items = list(dict.fromkeys(i for i in items if i))
    if len(items) < 2:
        raise ToolError("split_work is for a request that covers two or more separate documents; for one, read it "
                        "and call submit_plan")
    if len(items) > 10:
        raise ToolError("at most 10 items in one request")
    st.work_items = [{"request": i, "run_id": "", "status": "todo"} for i in items]
    ctx.log("work.split", items=items)
    return Result(f"Split into {len(items)} items, to be done one after another:\n"
                  + "\n".join(f"{n}. {i}" for n, i in enumerate(items, 1)), f"Split into {len(items)} items")


def submit_plan(ctx: Ctx, a: dict) -> Result:
    steps = [_clean_step(s) for s in a["steps"] if _clean_step(s)]
    if not 1 <= len(steps) <= 20:
        raise ToolError("give between 1 and 20 steps")
    st = ctx.state
    replan = st.phase == "act"
    kept = [s for s in st.plan if s.status != "todo"] if replan else []
    st.goal = str(a["goal"]).strip()
    new = [str(c).strip() for c in a.get("success_criteria") or [] if str(c).strip()]
    # the definition of done is fixed before acting: a re-plan may add to it, never drop it (seen 2026-10-09: a
    # re-plan sent no criteria, and the run was "verified" on its action read-backs alone)
    st.success_criteria = (st.success_criteria + [c for c in new if c not in st.success_criteria]) if replan else new
    st.plan = kept + [Step(t) for t in steps]
    st.phase = "act"
    ctx.log("plan.submit", goal=st.goal, success_criteria=st.success_criteria, steps=[s.text for s in st.plan],
            replan=replan)
    answers = []
    for q in [str(q).strip() for q in a.get("questions") or [] if str(q).strip()]:
        answers.append(ask_human(ctx, {"question": q}).text)
    nxt = st.next_step()
    text = (("Plan updated (open steps replaced)" if replan else "Plan accepted") + f". Goal: {st.goal}\n"
            + plan_text(st) + ("\n" + "\n".join(answers) if answers else "")
            + (f"\nNow do step {nxt}: {st.plan[nxt - 1].text}" if nxt else ""))
    return Result(text, f"Plan {'updated' if replan else 'accepted'}: {len(st.plan)} steps")


def step_done(ctx: Ctx, a: dict) -> Result:
    st = ctx.state
    if a.get("step") in (None, ""):  # no number given: the step shown as Next
        if not st.next_step():
            raise ToolError("every step is already marked; call finish")
        a = {**a, "step": st.next_step()}
    num = re.search(r"\d+", str(a["step"]))
    if not num:
        raise ToolError(f"step must be the number of a plan step (1–{len(st.plan)})")
    n = int(num[0])
    if not 1 <= n <= len(st.plan):
        raise ToolError(f"there is no step {n}; the plan has steps 1–{len(st.plan)}")
    result = str(a.get("result") or "").strip()
    step = st.plan[n - 1]
    if step.status != "todo" and st.open_steps():
        listing = "; ".join(f"{i}. {st.plan[i - 1].text}" for i in st.open_steps())
        raise ToolError(f"step {n} ({step.text}) was already marked {step.status}. The open steps are: {listing}. "
                        "If you just finished one of those, call step_done with its number.")
    step.status = "skipped" if re.match(r"^(skipped|not applicable|n/?a)\b", result, re.I) else "done"
    step.result = result
    ctx.log("plan.step_done", step=n, text=step.text, status=step.status, result=result)
    nxt = st.next_step()
    tail = f"Next: step {nxt} — {st.plan[nxt - 1].text}" if nxt else "All steps are done: call finish."
    return Result(f"Step {n} marked {step.status}. {tail}", f"Step {n} {step.status}: {result[:150]}")


def finish(ctx: Ctx, a: dict) -> Result:
    st = ctx.state
    open_ = st.open_steps()
    if open_:
        listing = "; ".join(f"{i}. {st.plan[i - 1].text}" for i in open_)
        raise ToolError(f"these steps are still open: {listing}. If you already did one, call step_done for it with "
                        f"what you did (do not do it again). Otherwise do it now, or mark it with "
                        f"step_done(step, result='skipped: <reason>') if it does not apply.")
    st.summary, st.finished, st.phase = str(a["summary"]).strip(), True, "done"
    ctx.log("run.finish", summary=st.summary)
    return Result("Finished.", "Finished.")


# --------------------------------------------------------------------------- the registry

S = {"type": "string"}


def _p(desc: str, **schema) -> dict:
    return {"type": "string", "description": desc} | schema


def build_registry(policy: Policy | None = None) -> Registry:
    act, both = ("act",), ("understand", "act")
    kinds = policy.approvals if policy else {}
    approval_help = "; ".join(f"{k}: {', '.join(v.get('details') or [])}" for k, v in kinds.items())
    return Registry([
        Tool("handbook_search", "Search the company handbook (procedures, approval rules, templates).",
             {"query": _p("what to look for")}, ["query"], handbook_search),
        Tool("handbook_read", "Read a handbook document, or one section of it.",
             {"document": _p("document name or number, e.g. 01_AP_Invoice_Processing or 06"),
              "section": _p("optional section heading, e.g. 'Step 6'")}, ["document"], handbook_read),
        Tool("mail_list", "List emails in the AP mailbox, newest first.",
             {"folder": _p("inbox (default), archive, sent or all"),
              "search": _p("optional text to find in the sender, subject or body")}, [], mail_list),
        Tool("mail_read", "Read one email: headers, body and its attachments.",
             {"id": {"type": "integer", "description": "message id, e.g. 9"}}, ["id"], mail_read),
        Tool("mail_attachment", "Open an email attachment and return its text (scans are read by OCR).",
             {"message_id": {"type": "integer"}, "attachment": _p("attachment id or file name")},
             ["message_id", "attachment"], mail_attachment),
        Tool("mail_send", "Send an email from the AP mailbox.",
             {"to": {"type": "array", "items": S, "description": "recipient addresses"},
              "subject": S, "body": _p("plain text"),
              "cc": {"type": "array", "items": S, "description": "optional cc addresses"},
              "attachments": {"type": "array", "items": S, "description": "optional drive file paths to attach"}},
             ["to", "subject", "body"], mail_send, kind="write", phases=act),
        Tool("mail_forward", "Forward an email, with its attachments, adding your message on top.",
             {"message_id": {"type": "integer"}, "to": {"type": "array", "items": S}, "body": S,
              "cc": {"type": "array", "items": S}}, ["message_id", "to", "body"], mail_forward, kind="write",
             phases=act),
        Tool("files_list", "List a folder on the shared drive (e.g. AP, Reports). Empty = drive root.",
             {"folder": _p("folder path relative to the drive root")}, [], files_list),
        Tool("file_read", "Read a file on the shared drive: PDF, spreadsheet (column names + last rows), CSV or text.",
             {"path": _p("path relative to the drive root, e.g. AP/AP_Tracker.xlsx"),
              "sheet": _p("optional sheet name"), "rows": {"type": "integer", "description": "spreadsheet rows to show (default 10)"}},
             ["path"], file_read),
        Tool("sheet_append_row", "Add one complete row at the bottom of a spreadsheet (e.g. the AP tracker): every "
             "column in one call. Existing rows are never changed.",
             {"path": _p("spreadsheet path, e.g. AP/AP_Tracker.xlsx"),
              "row": {"type": "object", "description": "{column name: value}, with the sheet's exact column names"},
              "sheet": _p("optional sheet name"),
              "blank": {"type": "array", "items": S,
                        "description": "only if the tool asks: columns left empty on purpose"},
              "correction": {"type": "integer",
                             "description": "only for a correction row: the number of the row it corrects"}},
             ["path", "row"], sheet_append_row, kind="write", phases=act),
        Tool("sheet_write", "Create a new spreadsheet file with one sheet.",
             {"path": _p("new file path ending in .xlsx"), "sheet": _p("sheet name"),
              "columns": {"type": "array", "items": S},
              "rows": {"type": "array", "items": {"type": "array"}, "description": "one list of values per row"}},
             ["path", "columns", "rows"], sheet_write, kind="write", phases=act),
        Tool("browser_open", "Open a page of a web app (the Finance Portal) and get a text digest of it.",
             {"url": _p("full URL or a path such as /bills or /vendors?q=Globex")}, ["url"], browser_open,
             phases=act),
        Tool("browser_fill", "Fill fields of the form on the current page, and optionally press a button to submit "
             "it. Returns the page that follows.",
             {"fields": {"type": "object", "description": "{field label: value} using the labels in the page digest; "
                                                          "for drop-downs, the option text or part of it"},
              "submit": _p("optional: the text of the button to press, e.g. 'Save bill'")},
             ["fields"], browser_fill, kind="write", phases=act),
        Tool("browser_click", "Click a link or button on the current page by its visible text, and get the page that "
             "follows.", {"target": _p("the visible text of the link or button")}, ["target"], browser_click,
             kind="write", phases=act),
        Tool("calc", "Exact arithmetic, e.g. '150000 * 0.02' or '177000 - 3000'. Use it for every calculation.",
             {"expression": S}, ["expression"], calc),
        Tool("date_calc", "Convert a date to DD-MM-YYYY, optionally add days, and compare it with today.",
             {"date": _p("a date in any common format, e.g. 30-Sep-2026"),
              "add_days": {"type": "integer", "description": "days to add (default 0)"}}, ["date"], date_calc),
        Tool("ask_human", "Ask the requester a question when the request is unclear, documents disagree, or the "
             "handbook says to ask. Waits for the answer.", {"question": S}, ["question"], ask_human, kind="human"),
        Tool("request_approval", "Ask the approver the Approval Matrix names to approve one specific action before "
             "you do it. Waits for the decision; returns an approval reference.",
             {"approver": _p("who must approve, e.g. CFO"),
              "action": _p("what kind of action", enum=list(kinds)),
              "details": {"type": "object", "description": f"exactly what is approved. Per action: {approval_help}"},
              "reason": _p("optional: why approval is needed")},
             ["action", "details"], request_approval, kind="human", phases=act),
        Tool("split_work", "Only when the request covers several separate documents (for example several invoices "
             "in the inbox): list one item per document. Each becomes its own task, done one after another.",
             {"items": {"type": "array", "items": S,
                        "description": "one self-contained request per document: what to handle and where it is, "
                                       "e.g. \"Process the invoice in email #12 (Vandelay, VAN-88231)\". Not what "
                                       "the outcome should be: that is decided when the item is done"}},
             ["items"], split_work, kind="control", phases=("understand",)),
        Tool("submit_plan", "When you have read the documents the request is about: give the goal and the steps of "
             "the handbook procedure you will follow, in order. Later, call it again to change the remaining steps.",
             {"goal": S, "success_criteria": {"type": "array", "items": S},
              "steps": {"type": "array", "items": S, "description": "short imperative steps, in order"},
              "questions": {"type": "array", "items": S, "description": "questions for the requester, if any"}},
             ["goal", "steps"], submit_plan, kind="control"),
        Tool("step_done", "Mark a plan step finished and note the facts you found or created (IDs, amounts, dates). "
             "The note stays visible to you.",
             {"step": {"type": "integer", "description": "the step number (default: the step shown as Next)"},
              "result": _p("what was done and found")}, ["result"], step_done,
             kind="control", phases=act),
        Tool("finish", "End the task once every plan step is done.", {"summary": _p("what was done, with IDs")},
             ["summary"], finish, kind="control", phases=act),
    ])
