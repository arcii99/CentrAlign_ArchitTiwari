"""VERIFY: an independent, read-only check of a finished run against the systems of record.

Nova saying "done" proves nothing. After finish, the verifier opens its own fresh session (a separate browser that
can only sign in and read, the mail API, the files on the drive) and checks two things:

1. Every action stuck. Each record Nova created is opened again and its fields compared with what was submitted;
   each email is looked up in the Sent folder; each tracker row is read back from the file.
2. The contract is met. Each success criterion written before acting is checked with a probe: a portal page and the
   values it must show, an email in the Sent folder, or a spreadsheet row. The model only chooses the probe for each
   criterion (it knows which page or row the criterion is about); reading and comparing is done by code.

A criterion no probe can check is reported as such, never as passed. And only a contradiction fails: a value the
page doesn't show is "unverified", because a wrong "failed" is not harmless. It sends Nova back to "fix" work that
was right (seen 2026-10-07: a check that matched "Status" to the page's "Status reason" got a paid bill reopened).
"""

from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook

from agent.browser import Browser
from agent.provenance import EMAIL, _dates, as_date, as_number
from agent.tools import ToolError, resolve_path

PROBES = {"type": "object", "properties": {"checks": {"type": "array", "items": {
    "type": "object",
    "properties": {
        "criterion": {"type": "integer"},
        "kind": {"type": "string", "enum": ["portal_page", "mail_sent", "sheet_row", "none"]},
        "path": {"type": "string"},
        "expect": {"type": "array", "items": {"type": "object", "properties": {
            "label": {"type": "string"}, "value": {"type": "string"}}, "required": ["label", "value"]}},
        "to": {"type": "string"},
        "subject_contains": {"type": "string"},
        "body_contains": {"type": "array", "items": {"type": "string"}},
    },
    "required": ["criterion", "kind"]}}}, "required": ["checks"]}

SYSTEM = ("You plan read-only checks of finished accounts-payable work. For each success criterion you choose where "
          "the company's systems would show it, using only the run record you are given. You never judge the result "
          "yourself; code reads the systems and compares.")

PROMPT = """Success criteria (written before the work started):
{criteria}

Run record:
{record}

For each criterion, choose one probe:
- portal_page: path (a Finance Portal page from the run record, e.g. /bills/BILL-2026-0001) and expect: the labelled values the page must show.
- mail_sent: to, subject_contains, and body_contains (short values the email body must contain, such as an invoice number or an amount).
- sheet_row: path (a spreadsheet on the drive, relative to the drive, e.g. AP/AP_Tracker.xlsx) and expect: column names and the values the row must have.
- none: when nothing in the systems can show it (for example a check that was only a calculation, or that something did not exist before).
Expect only what should still be true now that all the work is done: identifiers, amounts, dates, names, recipients. A status can move on as work continues (a bill that was entered and then paid shows Paid, not Posted), so expect a status only when the criterion is about the final status. Use the IDs, paths and values from the run record. Answer as JSON {{"checks": [...]}}, one per criterion."""


def _tokens(s: str) -> list[str]:
    return re.findall(r"[a-z0-9]+", str(s).lower())


def same(expected, observed) -> bool:
    """Is the observed value the expected one? Dates and amounts by meaning; text by its words."""
    e, o = str(expected).strip(), str(observed).strip()
    if as_date(e) and as_date(o):
        return as_date(e) == as_date(o)
    if as_number(e) is not None and as_number(o) is not None:
        return as_number(e) == as_number(o)
    words = _tokens(e)
    if bool(words) and all(w in _tokens(o) for w in words):
        return True
    # the same numbers, one said more briefly: "30 days" is "30 days from invoice date" (seen 2026-10-08)
    number = r"(?<![\w.])\d[\d,]*(?:\.\d+)?"
    ne, no = {as_number(n) for n in re.findall(number, e)}, {as_number(n) for n in re.findall(number, o)}
    we, wo = set(re.findall(r"[a-z]+", e.lower())), set(re.findall(r"[a-z]+", o.lower()))
    return bool(ne) and ne == no and (wo <= we or we <= wo)


def in_text(expected: str, text: str) -> bool:
    """Does a text (an email body, a page) contain the value? Amounts and dates by meaning."""
    e = str(expected).strip()
    if as_number(e) is not None:
        nums = {as_number(n) for n in re.findall(r"(?<![\w.])\d[\d,]*(?:\.\d+)?", text)}
        return as_number(e) in nums
    if as_date(e):
        return as_date(e) in {iso for _, iso in _dates(text)}
    return all(w in _tokens(text) for w in _tokens(e))


def values_in(phrase: str) -> list[str]:
    """The values a phrase carries ("Invoice No VAN-88231" → ["VAN-88231"]; "Total 70,800.00" → ["70,800.00"]): an
    email says them in its own words (seen 2026-10-09: "no email … with ['Invoice No VAN-88231']" for one that had
    "Invoice number: VAN-88231"). A bare value is itself."""
    p = str(phrase).strip()
    if " " not in p or as_number(p) is not None or as_date(p):
        return [p]
    found = [raw for raw, _ in _dates(p)]
    rest = p
    for raw in found:
        rest = rest.replace(raw, " ")
    found += re.findall(r"(?<![\w.])\d[\d,]*(?:\.\d+)?(?![\w/-])|(?<![\w/.-])[A-Za-z0-9][A-Za-z0-9/._-]*\d[A-Za-z0-9/._-]*",
                        rest)  # amounts first: "70,800.00" whole, not "70" and "800.00"
    return found or [p]


def plain_value(v) -> str:
    """A small model sometimes writes {"amount": 150000.0, "currency": "INR"} where a value belongs."""
    if isinstance(v, str) and v.strip().startswith("{"):
        try:
            v = json.loads(v)
        except ValueError:
            return v.strip()
    if isinstance(v, dict):
        for k in ("amount", "value", "text", "name", "date"):
            if v.get(k) not in (None, ""):
                return plain_value(v[k])
        scalars = [x for x in v.values() if isinstance(x, (str, int, float)) and not isinstance(x, bool)]
        return plain_value(scalars[0]) if scalars else ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return "" if v is None else str(v).strip()


def grounded(value: str, criterion: str, record: str = "") -> bool:
    """Did an expected value come from the definition of done, rather than from the planner's imagination? Amounts,
    dates and IDs may come from the criterion or the run record (they are blanked out of criteria). Words (a status,
    a name) only from the criterion: the record says what Nova did, and comparing the systems with that is the
    action read-backs' job. (Seen 2026-10-07/08: the planner expected a tracker row's status to be "Paid", once with
    nothing behind it, once "grounded" by a stray "paid" somewhere in the record.)"""
    text = f"{criterion}\n{record}"
    if re.search(r"\d", value):  # amounts, dates and identifiers: by meaning
        if in_text(value, text):
            return True
        # mixed ("TDS: ₹3,000"): every number by its value, every word as a word
        number = r"(?<![\w.])\d[\d,]*(?:\.\d+)?"
        known = {as_number(n) for n in re.findall(number, text)}
        return all(as_number(n) in known for n in re.findall(number, value)) and \
            all(w in _tokens(text) for w in re.findall(r"[a-z]+", value.lower()))
    phrase = re.sub(r"\s+", " ", value.strip())
    return bool(phrase) and re.search(r"(?<!\w)" + re.escape(phrase) + r"(?!\w)", re.sub(r"\s+", " ", criterion),
                                      re.I) is not None


def _rows(path: Path) -> list[list]:
    """All rows of a workbook's first sheet. Closed straight away: a read-only workbook keeps the file open, and on
    Windows an open file can't be replaced or saved by anyone else."""
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    finally:
        wb.close()


def _pairs(snap: dict | None) -> list[tuple[str, str]]:
    return [tuple(p) for t in (snap or {}).get("tables", []) if t["kind"] == "kv" for p in t["pairs"]]


def record_text(snap: dict | None) -> str:
    """What a page shows as data: its heading and its tables. Not its forms: a status form's options list every
    status there is, so searching them would find any status at all."""
    snap = snap or {}
    parts = [snap.get("h1") or "", snap.get("title") or ""]
    for t in snap.get("tables", []):
        if t["kind"] == "kv":
            parts += [f"{k}: {v}" for k, v in t["pairs"]]
        else:
            parts += [" | ".join(t.get("headers") or [])] + [" | ".join(r) for r in t.get("rows") or []]
    return "\n".join(parts)


_SAME_WORD = {"reference": "ref", "number": "no", "num": "no", "qty": "quantity"}
_GENERIC = {"amount", "no", "value", "id", "ref"}  # "TDS amount" is "TDS", "PO number" is "PO"


def _label_words(label: str) -> list[str]:
    words = [_SAME_WORD.get(w, w) for w in _tokens(label)]
    while len(words) > 1 and words[-1] in _GENERIC:
        words = words[:-1]
    return words


def label_match(want: str, have: str) -> str:
    """How a page's label relates to the one asked for: "exact" (same words, ignoring a generic last word like
    amount/number), "close" (one ends with the other: "Method" for "Payment method") or "" (different: "Status
    reason" is not "Status")."""
    w, h = _label_words(want), _label_words(have)
    if not w or not h:
        return ""
    if w == h:
        return "exact"
    return "close" if w[-len(h):] == h or h[-len(w):] == w else ""


def check_value(label: str, value, pairs: list[tuple[str, str]], text: str) -> tuple[str, str]:
    """("ok" | "contradicts" | "missing", what was seen). The same label with another value contradicts; a value shown
    anywhere on the page (a heading, under another label) is ok; a value the page doesn't show is only missing."""
    exact = [v for k, v in pairs if label_match(label, k) == "exact"]
    if exact:
        return ("ok", exact[0]) if any(same(value, v) for v in exact) else ("contradicts", exact[0])
    if in_text(value, text):
        return "ok", "shown on the page"
    close = [(k, v) for k, v in pairs if label_match(label, k) == "close"]
    if len(close) == 1:  # one obvious counterpart ("Method" for "Payment method"): treat it as the same field
        k, v = close[0]
        return ("ok", v) if same(value, v) else ("contradicts", f"{v} (as {k})")
    return "missing", ""


class Verifier:
    def __init__(self, sites: dict, mail, drive: Path, run_dir: Path, llm):
        self.mail, self.drive, self.run_dir, self.llm = mail, drive, run_dir, llm
        self.browser = Browser(sites, run_dir / "verify", headless=True, on_post=self._read_only)
        self.records: set[str] = set()  # the pages of the records this run created or changed

    @staticmethod
    def _read_only(req: dict):
        """The verifier's browser can sign in and read; nothing else it sends would get through."""
        return None if "******" in (req.get("form") or {}).values() else (False, "the verifier only reads")

    def close(self) -> None:
        self.browser.close()

    # ------------------------------------------------------------------ 1. did every action stick?
    def actions(self, events: list[dict]) -> list[dict]:
        checks, call, post = [], {}, None
        for e in events:
            kind, d = e["kind"], e["data"]
            if kind == "tool.call":
                call, post = d, None
            elif kind == "browser.post" and d.get("action"):
                post = d
            elif kind == "tool.result" and post and d["ok"] and d["name"] in ("browser_fill", "browser_click"):
                fields = (call.get("arguments") or {}).get("fields") or {}
                checks.append(self._record(post["action"], str(d["data"].get("url", "")).split("?")[0], fields))
                post = None
            elif kind == "mail.sent":
                checks.append(self._sent(d))
            elif kind == "file.row_appended":
                checks.append(self._row(d))
        return checks

    def _record(self, action: str, url: str, fields: dict) -> dict:
        path = "/" + url.split("://", 1)[-1].split("/", 1)[-1]
        what = f"{action} → {path}"
        self.records.add(path)
        res, snap = self.browser.inspect(url)
        if not res.ok or snap is None:
            return {"what": what, "status": "fail", "detail": f"the record could not be opened: {res.summary[:200]}"}
        pairs, good, bad = _pairs(snap), [], []
        for label, value in fields.items():
            if str(value).strip() == "":
                continue
            verdict, seen = check_value(label, value, pairs, record_text(snap))
            if verdict == "ok":
                good.append(label)
            elif verdict == "contradicts":
                bad.append(f"{label}: submitted {value}, record shows {seen}")
        detail = (f"{len(good)} submitted value(s) confirmed on the record" + (f"; mismatches: {'; '.join(bad)}"
                                                                               if bad else ""))
        return {"what": what, "status": "fail" if bad else "pass", "detail": detail, "observed": dict(pairs)}

    def _sent(self, d: dict) -> dict:
        what = f"email #{d['id']} to {', '.join(d['to'])}: {d['subject']}"
        try:
            m = self.mail.read(int(d["id"]))
        except Exception as e:  # not there, or the mail server is down: either way it isn't confirmed
            return {"what": what, "status": "fail", "detail": f"not found in the mailbox ({e})"}
        ok = m["folder"] == "sent" and m["subject"] == d["subject"] and all(a in m["to"] for a in d["to"])
        return {"what": what, "status": "pass" if ok else "fail",
                "detail": f"in Sent, to {m['to']}" if ok else f"found in {m['folder']} to {m['to']}: {m['subject']}"}

    def _row(self, d: dict) -> dict:
        what = f"row {d['row']} of {d['path']}"
        try:
            rows = _rows(self._file(d["path"]))
            header, row = rows[0], rows[d["row"] - 1]
        except Exception as e:
            return {"what": what, "status": "fail", "detail": f"could not read it back ({e})"}
        found = {str(h): ("" if v is None else str(v)) for h, v in zip(header, row) if h}
        bad = [f"{k}: written {v}, file has {found.get(k)}" for k, v in d["values"].items()
               if v and not same(v, found.get(k, ""))]
        return {"what": what, "status": "fail" if bad else "pass",
                "detail": "; ".join(bad) or "every value read back from the file"}

    # ------------------------------------------------------------------ 2. is the contract met?
    def criteria(self, criteria: list[str], record: str) -> list[dict]:
        if not criteria:
            return []
        text = "\n".join(f"{i}. {c}" for i, c in enumerate(criteria, 1))
        try:
            reply = self.llm.chat([{"role": "system", "content": SYSTEM},
                                   {"role": "user", "content": PROMPT.format(criteria=text, record=record)}],
                                  schema=PROBES, think=False)
            probes = {p.get("criterion"): p for p in json.loads(reply.content or "{}").get("checks", [])
                      if isinstance(p, dict)}
        except Exception as e:  # no plan for the probes: every criterion stays unverified, and says why
            probes = {}
            reason = f"no probe could be planned ({e.__class__.__name__})"
        else:
            reason = "no probe was chosen for it"
        out = []
        for i, c in enumerate(criteria, 1):
            p = probes.get(i) or {"kind": "none"}
            check = {"criterion": c, "probe": p}
            if p["kind"] == "none" or not p.get("path") and p["kind"] != "mail_sent":
                check.update(status="unverified", detail=reason if i not in probes else "nothing in the systems shows it")
            else:
                check.update(self._probe(p, (c, record)))
            out.append(check)
        return out

    def _file(self, raw: str) -> Path:
        """A drive path from the planner, resolved like Nova's own file tools: never outside the drive."""
        return resolve_path(self.drive, self.run_dir, raw)

    def _probe(self, p: dict, grounds: tuple[str, str] | None = None) -> dict:
        """One probe. `grounds` is (the criterion, the run record): an expected value they don't support is the
        planner's guess, which can't confirm or contradict anything, so it is set aside (and named in the result)."""
        raw = [(str(e.get("label", "")), plain_value(e.get("value", ""))) for e in p.get("expect") or []
               if isinstance(e, dict)]
        raw = [(lbl, v) for lbl, v in raw if v]
        expect = [(lbl, v) for lbl, v in raw if not grounds or grounded(v, *grounds)]
        guessed = [f"{lbl}={v}" for lbl, v in raw if (lbl, v) not in expect]
        result = self._check(p, expect, grounds)
        if guessed:
            result["detail"] += f" (set aside, not from the criterion or the record: {', '.join(guessed)})"
        return result

    def _check(self, p: dict, expect: list[tuple[str, str]], grounds: tuple[str, str] | None) -> dict:
        if p["kind"] == "portal_page":
            res, snap = self.browser.inspect(p["path"])
            if not res.ok or snap is None:
                # a record this run made that won't open is a failure; any other page may just be a wrong guess
                mine = str(p["path"]).split("?")[0] in self.records
                return {"status": "fail" if mine else "unverified",
                        "detail": f"{p['path']} could not be opened: {res.summary[:160]}"}
            if not expect:
                return {"status": "unverified", "detail": f"{p['path']} exists; the probe named no values to check"}
            pairs, bad, missing = _pairs(snap), [], []
            for label, value in expect:
                verdict, seen = check_value(label, value, pairs, record_text(snap))
                if verdict == "contradicts":
                    bad.append(f"{label}: expected {value}, page shows {seen}")
                elif verdict == "missing":
                    missing.append(f"{label} {value}")
            if bad:
                return {"status": "fail", "detail": "; ".join(bad)}
            if missing:
                return {"status": "unverified", "detail": f"{p['path']} does not show: {'; '.join(missing)}"}
            return {"status": "pass", "detail": f"{p['path']} shows it"}
        if p["kind"] == "mail_sent":
            return self._mail_probe(p, expect, grounds)
        if p["kind"] == "sheet_row":
            return self._sheet_probe(p, expect)
        return {"status": "unverified", "detail": f"unknown probe {p['kind']!r}"}

    def _mail_probe(self, p: dict, expect: list[tuple[str, str]], grounds: tuple[str, str] | None = None) -> dict:
        # a small model sometimes puts the probe's own fields inside expect: {"label": "subject_contains", ...}
        given = {" ".join(_tokens(k)): v for k, v in expect}
        to = str(p.get("to") or given.get("to") or "").strip().lower()
        subj = str(p.get("subject_contains") or given.get("subject contains") or given.get("subject") or "").strip()
        body = [plain_value(b) for b in (p.get("body_contains") or [])]
        body += [v for k, v in given.items() if k in ("body contains", "body")]
        # one string listing several values ("GLX/2026/0917, 30-Sep-2026, 177000.00") is several values; an Indian
        # amount (1,77,000.00) has no space after its commas, so it stays whole
        body = [part.strip() for b in body for part in re.split(r",\s+|;\s*", b) if part.strip()]
        # values only: a description ("invoice number", "and method") says what to look for, not what is there
        body = [b for b in body if re.search(r"\d", b) or re.fullmatch(r"[A-Z][A-Z0-9/-]{2,}", b)]
        body = [v for b in body for v in values_in(b)]
        body = [b for b in body if not grounds or grounded(b, *grounds)]
        # whoever the criterion names must have received it, whatever the planner wrote (seen 2026-10-09: the planner
        # copied "anya.iyer@…" from Nova's own misdirected email, that guess was set aside, and with no recipient left
        # to check, the email passed for a criterion that names ananya.iyer@…)
        own = str(getattr(self.mail, "mailbox", "")).lower()
        named = [e.lower() for e in EMAIL.findall(grounds[0] if grounds else "") if e.lower() != own]
        if not (to or subj or body or named):
            return {"status": "unverified", "detail": "the probe named no recipient, subject or content to look for"}
        near = None
        for m in self.mail.list("sent", subj[:60] if subj else (to or (named[0] if named else "")), 50):
            text = self.mail.read(m["id"])["body"] + " " + m["subject"]
            recipients = (m["to"] + "," + m["cc"]).lower()
            if (not subj or in_text(subj, m["subject"])) and all(in_text(b, text) for b in body):
                if (not to or to in recipients) and all(n in recipients for n in named):
                    return {"status": "pass", "detail": f"email #{m['id']} to {m['to']}: {m['subject']}"}
                near = near or m
        if near is not None and named:
            return {"status": "fail", "detail": f"email #{near['id']} ({near['subject']}) went to {near['to']}"
                                                + (f", cc {near['cc']}" if near["cc"] else "")
                                                + f", but the criterion names {', '.join(named)}"}
        return {"status": "fail", "detail": f"no email in Sent to {to or ', '.join(named) or 'anyone'} matching "
                                            f"{subj!r} with {body}"}

    def _sheet_probe(self, p: dict, expect: list[tuple[str, str]]) -> dict:
        try:
            rows = _rows(self._file(p.get("path") or ""))
        except (ToolError, OSError, ValueError) as e:
            return {"status": "unverified", "detail": f"{p.get('path')} could not be read ({e})"}
        header = [str(h or "") for h in rows[0]] if rows else []
        cols = {lbl: [i for i, h in enumerate(header) if label_match(lbl, h) == "exact"] for lbl, _ in expect}
        usable = [(lbl, v) for lbl, v in expect if cols[lbl]]
        if not usable:
            return {"status": "unverified", "detail": f"{p.get('path')} has no column like "
                                                      + ", ".join(lbl for lbl, _ in expect)}
        best, best_n, best_bad = 0, -1, []
        for n, row in enumerate(rows[1:], 2):
            bad = [f"{lbl} is {row[cols[lbl][0]]!r}, expected {v!r}" for lbl, v in usable
                   if not any(same(v, "" if row[i] is None else row[i]) for i in cols[lbl])]
            if not bad:
                unused = [lbl for lbl, _ in expect if not cols[lbl]]
                return {"status": "unverified" if unused else "pass",
                        "detail": f"row {n} of {p['path']}" + (f" (no column for: {', '.join(unused)})" if unused else "")}
            if len(usable) - len(bad) > best_n:
                best, best_n, best_bad = n, len(usable) - len(bad), bad
        return {"status": "fail", "detail": f"no row in {p['path']} with " + ", ".join(f"{l}={v}" for l, v in usable)
                                            + (f"; closest is row {best}: {'; '.join(best_bad)}" if best_n > 0 else "")}

    # ------------------------------------------------------------------ the verdict
    def run(self, criteria: list[str], events: list[dict], record: str) -> dict:
        actions = self.actions(events)
        checks = self.criteria(criteria, record)
        statuses = [c["status"] for c in actions + checks]
        verdict = "failed" if "fail" in statuses else "verified" if "pass" in statuses else "unverified"
        if verdict == "verified" and not checks:  # read-backs say the actions were recorded, not that the request
            verdict = "unverified"                 # was met: without criteria nothing is verified
        return {"verdict": verdict, "actions": actions, "criteria": checks,
                "counts": {s: statuses.count(s) for s in ("pass", "fail", "unverified")}}
