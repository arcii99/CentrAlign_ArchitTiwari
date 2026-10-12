"""Generators for the company's documents: invoices (text and scanned), vendor onboarding papers, spreadsheets."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

# --------------------------------------------------------------------------- formatting


def inr(x: float) -> str:
    """Indian digit grouping: 177000 -> '1,77,000.00'."""
    sign = "-" if x < 0 else ""
    whole, frac = f"{abs(round(x, 2)):.2f}".split(".")
    if len(whole) > 3:
        head, tail = whole[:-3], whole[-3:]
        groups = []
        while len(head) > 2:
            groups.insert(0, head[-2:])
            head = head[:-2]
        if head:
            groups.insert(0, head)
        whole = ",".join([*groups, tail])
    return f"{sign}{whole}.{frac}"


_ONES = ("", "One", "Two", "Three", "Four", "Five", "Six", "Seven", "Eight", "Nine", "Ten", "Eleven", "Twelve",
         "Thirteen", "Fourteen", "Fifteen", "Sixteen", "Seventeen", "Eighteen", "Nineteen")
_TENS = ("", "", "Twenty", "Thirty", "Forty", "Fifty", "Sixty", "Seventy", "Eighty", "Ninety")


def _two(n: int) -> str:
    return _ONES[n] if n < 20 else _TENS[n // 10] + (f" {_ONES[n % 10]}" if n % 10 else "")


def _words(n: int) -> str:
    """Indian numbering system: crore, lakh, thousand."""
    if n == 0:
        return "Zero"
    crore, n = divmod(n, 10**7)
    lakh, n = divmod(n, 10**5)
    thousand, n = divmod(n, 1000)
    hundred, n = divmod(n, 100)
    parts = []
    if crore:
        parts.append(f"{_words(crore)} Crore")
    if lakh:
        parts.append(f"{_two(lakh)} Lakh")
    if thousand:
        parts.append(f"{_two(thousand)} Thousand")
    if hundred:
        parts.append(f"{_ONES[hundred]} Hundred")
    if n:
        parts.append(_two(n))
    return " ".join(parts)


def amount_in_words(x: float) -> str:
    rupees = int(round(x, 2))
    paise = int(round((x - rupees) * 100))
    text = f"Rupees {_words(rupees)}"
    if paise:
        text += f" and {_words(paise)} Paise"
    return text + " Only"


# --------------------------------------------------------------------------- invoices

_styles = getSampleStyleSheet()
_small = ParagraphStyle("small", parent=_styles["Normal"], fontSize=8.5, leading=11)
_normal = ParagraphStyle("norm", parent=_styles["Normal"], fontSize=9.5, leading=12.5)
_title = ParagraphStyle("title", parent=_styles["Title"], fontSize=15, spaceAfter=4)
_vendor = ParagraphStyle("vendor", parent=_styles["Heading1"], fontSize=17, spaceAfter=2)


def invoice_amounts(lines: list[dict], gst_rate: float) -> tuple[float, float, float]:
    taxable = round(sum(l["qty"] * l["rate"] for l in lines), 2)
    gst = round(taxable * gst_rate / 100, 2)
    return taxable, gst, round(taxable + gst, 2)


def render_invoice(path: Path, inv: dict) -> None:
    """inv keys: vendor{legal_name,address,gstin,pan,email,phone,state_code}, buyer{...}, invoice_no,
    invoice_date (date), date_fmt, po_ref, lines[{item,hsn_sac,qty,unit,rate}], gst_rate, bank{...}, note."""
    v, b = inv["vendor"], inv["buyer"]
    taxable, gst, total = invoice_amounts(inv["lines"], inv["gst_rate"])
    inter_state = v["state_code"] != b["state_code"]
    story = [
        Paragraph(v["legal_name"], _vendor),
        Paragraph(f"{v['address']}<br/>GSTIN: {v['gstin']} &nbsp;|&nbsp; PAN: {v['pan']}<br/>"
                  f"{v['email']} &nbsp;|&nbsp; {v.get('phone', '')}", _small),
        Spacer(1, 6 * mm),
        Paragraph("TAX INVOICE", _title),
    ]
    meta = [
        ["Invoice No.", inv["invoice_no"], "Bill To", ""],
        ["Invoice Date", inv["invoice_date"].strftime(inv.get("date_fmt", "%d-%b-%Y")),
         Paragraph(f"<b>{b['legal_name']}</b><br/>{b['address']}<br/>GSTIN: {b['gstin']}", _small), ""],
        ["PO Reference", inv.get("po_ref") or "—", "", ""],
        ["Place of Supply", f"Karnataka ({b['state_code']})", "", ""],
    ]
    if inv.get("period"):
        meta.append(["Service Period", inv["period"], "", ""])
    t = Table(meta, colWidths=[30 * mm, 50 * mm, 22 * mm, 75 * mm])
    t.setStyle(TableStyle([
        ("FONT", (0, 0), (-1, -1), "Helvetica", 9),
        ("FONT", (0, 0), (0, -1), "Helvetica-Bold", 9),
        ("FONT", (2, 0), (2, 0), "Helvetica-Bold", 9),
        ("SPAN", (2, 1), (3, len(meta) - 1)),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story += [t, Spacer(1, 5 * mm)]

    rows = [["#", "Description", "HSN/SAC", "Qty", "Unit", "Rate (INR)", "Amount (INR)"]]
    for i, l in enumerate(inv["lines"], 1):
        rows.append([str(i), Paragraph(l["item"], _small), l["hsn_sac"], f"{l['qty']:g}", l["unit"],
                     inr(l["rate"]), inr(l["qty"] * l["rate"])])
    rows.append(["", "Taxable Value", "", "", "", "", inr(taxable)])
    if inter_state:
        rows.append(["", f"IGST @ {inv['gst_rate']:g}%", "", "", "", "", inr(gst)])
    else:
        half = inv["gst_rate"] / 2
        rows.append(["", f"CGST @ {half:g}%", "", "", "", "", inr(gst / 2)])
        rows.append(["", f"SGST @ {half:g}%", "", "", "", "", inr(gst / 2)])
    rows.append(["", "Invoice Total", "", "", "", "", inr(total)])
    lt = Table(rows, colWidths=[8 * mm, 70 * mm, 20 * mm, 14 * mm, 14 * mm, 25 * mm, 28 * mm])
    lt.setStyle(TableStyle([
        ("FONT", (0, 0), (-1, -1), "Helvetica", 9),
        ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 9),
        ("FONT", (0, -1), (-1, -1), "Helvetica-Bold", 9.5),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8e8e8")),
        ("GRID", (0, 0), (-1, len(inv["lines"])), 0.4, colors.grey),
        ("LINEABOVE", (0, -1), (-1, -1), 0.8, colors.black),
        ("ALIGN", (3, 0), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story += [lt, Spacer(1, 3 * mm),
              Paragraph(f"<b>Amount in words:</b> {amount_in_words(total)}", _normal), Spacer(1, 5 * mm)]
    bank = inv.get("bank")
    if bank:
        story.append(Paragraph(
            f"<b>Bank details for payment</b><br/>Account name: {bank['holder_name']}<br/>"
            f"Bank: {bank['bank_name']}, {bank.get('branch', '')}<br/>"
            f"A/c No: {bank['account_number']} &nbsp;|&nbsp; IFSC: {bank['ifsc']}", _small))
        story.append(Spacer(1, 4 * mm))
    if inv.get("note"):
        story += [Paragraph(inv["note"], _small), Spacer(1, 4 * mm)]
    story += [Spacer(1, 10 * mm), Paragraph(f"For {v['legal_name']}<br/><br/>Authorised Signatory", _small),
              Spacer(1, 6 * mm), Paragraph("This is a computer-generated invoice.", _small)]
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=14 * mm,
                      bottomMargin=14 * mm, title=f"Invoice {inv['invoice_no']}").build(story)


def rasterize(src: Path, dst: Path, angle: float = 0.6, dpi: int = 130) -> None:
    """Turn a PDF into an image-only 'scanned' PDF: no extractable text, slight skew, grain and blur."""
    import pypdfium2 as pdfium
    from PIL import Image, ImageFilter

    pdf = pdfium.PdfDocument(str(src))
    pages = []
    for page in pdf:
        img = page.render(scale=dpi / 72).to_pil().convert("L")
        img = img.rotate(angle, resample=Image.BICUBIC, expand=True, fillcolor=255)
        img = Image.blend(img, Image.effect_noise(img.size, 22), 0.07)
        pages.append(img.filter(ImageFilter.GaussianBlur(0.55)).convert("RGB"))
    pdf.close()
    pages[0].save(dst, "PDF", resolution=dpi, save_all=True, append_images=pages[1:])


# --------------------------------------------------------------------------- vendor onboarding papers


def _kv_table(rows: list[tuple[str, str]], widths=(60 * mm, 110 * mm)) -> Table:
    t = Table([[k, Paragraph(v, _normal)] for k, v in rows], colWidths=list(widths))
    t.setStyle(TableStyle([
        ("FONT", (0, 0), (0, -1), "Helvetica-Bold", 9.5),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f2f2f2")),
    ]))
    return t


def render_registration_form(path: Path, f: dict) -> None:
    h2 = _styles["Heading3"]
    story = [
        Paragraph("Northwind Robotics Pvt. Ltd.", _small),
        Paragraph("VENDOR REGISTRATION FORM", _title),
        Paragraph("To be completed by the vendor. Attach: GST registration certificate and a cancelled cheque or "
                  "bank letter.", _small),
        Spacer(1, 4 * mm), Paragraph("1. Business details", h2),
        _kv_table([("Legal name", f["legal_name"]), ("Trade name", f["trade_name"]),
                   ("Constitution", f["constitution"]), ("GSTIN", f["gstin"]), ("PAN", f["pan"]),
                   ("State", f["state"]), ("Registered address", f["address"]),
                   ("Nature of supply", f["nature"])]),
        Spacer(1, 4 * mm), Paragraph("2. Contact for accounts", h2),
        _kv_table([("Contact person", f["contact_name"]), ("Email", f["email"]), ("Phone", f["phone"])]),
        Spacer(1, 4 * mm), Paragraph("3. Bank details for payments", h2),
        _kv_table([("Account holder", f["bank"]["holder_name"]), ("Account number", f["bank"]["account_number"]),
                   ("IFSC", f["bank"]["ifsc"]), ("Bank & branch", f"{f['bank']['bank_name']}, {f['bank']['branch']}")]),
        Spacer(1, 4 * mm), Paragraph("4. Commercial terms", h2),
        _kv_table([("Payment terms", f["terms"])]),
        Spacer(1, 6 * mm),
        Paragraph("I confirm that the information above is correct.", _normal), Spacer(1, 8 * mm),
        Paragraph(f"Signed: <i>{f['contact_name']}</i>, {f['signatory_title']} &nbsp;&nbsp; Date: {f['signed_on']}",
                  _normal),
    ]
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
                      bottomMargin=16 * mm, title="Vendor Registration Form").build(story)


def render_gst_certificate(path: Path, f: dict) -> None:
    story = [
        Paragraph("Certificate of GST Registration", _title),
        Paragraph("(specimen issued for the Northwind Robotics mock environment)", _small),
        Spacer(1, 6 * mm),
        _kv_table([("Registration Number (GSTIN)", f["gstin"]), ("Legal Name", f["legal_name"]),
                   ("Trade Name", f["trade_name"]), ("Constitution of Business", f["constitution"]),
                   ("Principal Place of Business", f["address"]), ("Date of Liability", f["liability_date"]),
                   ("Type of Registration", "Regular")]),
    ]
    SimpleDocTemplate(str(path), pagesize=A4, title="GST Registration Certificate").build(story)


def render_cancelled_cheque(path: Path, c: dict) -> None:
    """Draw a CTS-style cheque, cross it out as CANCELLED, then 'scan' it (image only)."""
    raw = path.with_suffix(".raw.pdf")
    w, h = 203 * mm, 92 * mm
    cv = canvas.Canvas(str(raw), pagesize=(w, h))
    cv.setFillColor(colors.HexColor("#eef3f8"))
    cv.rect(0, 0, w, h, stroke=0, fill=1)
    cv.setFillColor(colors.HexColor("#123a6b"))
    cv.setFont("Helvetica-Bold", 15)
    cv.drawString(8 * mm, h - 12 * mm, c["bank_name"].upper())
    cv.setFont("Helvetica", 7.5)
    cv.setFillColor(colors.black)
    cv.drawString(8 * mm, h - 17 * mm, f"{c['branch']} Branch  |  IFSC: {c['ifsc']}")
    cv.drawString(w - 62 * mm, h - 12 * mm, "Date  D D M M Y Y Y Y")
    cv.setFont("Helvetica", 9)
    cv.drawString(8 * mm, h - 30 * mm, "Pay ________________________________________________ or Bearer")
    cv.drawString(8 * mm, h - 40 * mm, "Rupees ______________________________________________________")
    cv.rect(w - 52 * mm, h - 44 * mm, 44 * mm, 9 * mm)
    cv.drawString(w - 50 * mm, h - 41 * mm, "Rs.")
    cv.setFont("Helvetica-Bold", 10)
    cv.drawString(8 * mm, h - 54 * mm, f"A/c No. {c['account_number']}")
    cv.setFont("Helvetica", 8)
    cv.drawString(w - 85 * mm, h - 62 * mm, f"For {c['holder_name'].upper()}")
    cv.drawString(w - 70 * mm, h - 75 * mm, "Authorised Signatories")
    cv.setFont("Courier-Bold", 11)
    cv.drawString(30 * mm, 6 * mm, f"\"{c['cheque_no']}\"  {c['micr']}  {c['account_number'][-6:]}  29")
    cv.saveState()
    cv.translate(w / 2, h / 2)
    cv.rotate(12)
    cv.setStrokeColor(colors.HexColor("#333333"))
    cv.setLineWidth(1.4)
    cv.line(-55 * mm, 9 * mm, 80 * mm, 9 * mm)  # stop short of the account number, as a careful signer would
    cv.line(-55 * mm, -9 * mm, 80 * mm, -9 * mm)
    cv.setFont("Helvetica-Bold", 26)
    cv.drawCentredString(0, -4 * mm, "CANCELLED")
    cv.restoreState()
    cv.save()
    rasterize(raw, path, angle=-0.9, dpi=160)
    raw.unlink()


# --------------------------------------------------------------------------- spreadsheets

TRACKER_COLUMNS = ["Date Processed", "Vendor", "Vendor Invoice No", "Invoice Date", "PO No", "Invoice Total", "TDS",
                   "Net Payable", "Bill ID", "Payment ID", "Payment Date", "Status", "Notes"]
_HEADER_FILL = PatternFill("solid", fgColor="1F3B5C")


def _styled_header(ws, columns: list[str], widths: list[int]) -> None:
    ws.append(columns)
    for i, cell in enumerate(ws[1]):
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        ws.column_dimensions[cell.column_letter].width = widths[i]
    ws.freeze_panes = "A2"


def write_tracker(path: Path, rows: list[list]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoices"
    _styled_header(ws, TRACKER_COLUMNS, [14, 34, 18, 13, 14, 13, 9, 13, 15, 15, 13, 26, 48])
    for r in rows:
        ws.append(r)
    wb.save(path)


SPEND_COLUMNS = ["Vendor ID", "Vendor", "Payments", "Total Paid (INR)"]


def write_spend_report(path: Path, title: str, rows: list[tuple[str, str, int, float]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    _styled_header(ws, SPEND_COLUMNS, [12, 40, 11, 18])
    for r in sorted(rows, key=lambda r: -r[3]):
        ws.append(list(r))
    ws.append(["TOTAL", "", sum(r[2] for r in rows), round(sum(r[3] for r in rows), 2)])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        row[0].number_format = "#,##0.00"
    ws.sheet_properties.tabColor = "1F3B5C"
    wb.properties.title = title
    wb.save(path)


def dmy(d: date) -> str:
    return d.strftime("%d-%m-%Y")
